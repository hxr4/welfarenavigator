import json
import sys

import openpyxl

edits_path, workbook_path = sys.argv[1], sys.argv[2]
edits = json.load(open(edits_path, encoding="utf-8"))
wb = openpyxl.load_workbook(workbook_path)


def header(ws):
    return {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}


applied = 0
for e in edits["rows"]:
    ws = wb[e["sheet"]]
    h = header(ws)
    seen = -1
    for r in range(2, ws.max_row + 1):
        cell = lambda k: str(ws.cell(r, h[k]).value or "").strip()
        if cell("scheme_id") != e["scheme_id"]:
            continue
        if "condition_id" in e and cell("condition_id") != e["condition_id"]:
            continue
        if "doc_id" in e and cell("doc_id") != e["doc_id"]:
            continue
        if "index" in e:
            seen += 1
            if seen != e["index"]:
                continue
        for k, v in e["set"].items():
            if k in h:
                ws.cell(r, h[k]).value = v
        applied += 1
        break

ws = wb["sources"]
h = header(ws)
for r in range(2, ws.max_row + 1):
    if str(ws.cell(r, h["source_id"]).value or "").strip() in edits["sources"]:
        ws.cell(r, h["accessed_on"]).value = edits["date"]

wb.save(workbook_path)
print(json.dumps({"applied": applied}))
