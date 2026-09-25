import os
import openpyxl

P = "dataset/welfare-dataset.xlsx"
wb = openpyxl.load_workbook(P)
TODAY = "2026-09-23"
PR = "PENDING_SECOND_REVIEW"
K = "SRC-KFWFB-001"
ACT = "SRC-PLNT-001"


def col(ws, name):
    return [c.value for c in ws[1]].index(name) + 1


def find(ws, **kv):
    return [r for r in range(2, ws.max_row + 1) if all(ws.cell(r, col(ws, k)).value == v for k, v in kv.items())]


def append(ws, d):
    r = ws.max_row + 1
    while r > 2 and all(ws.cell(r - 1, c).value is None for c in range(1, ws.max_column + 1)):
        r -= 1
    for k, v in d.items():
        ws.cell(r, col(ws, k)).value = v


def setv(ws, r, k, v):
    ws.cell(r, col(ws, k)).value = v


def get(ws, r, k):
    return ws.cell(r, col(ws, k)).value


def upsert(ws, key, d):
    rows = find(ws, **{key: d[key]})
    if rows:
        for k, v in d.items():
            setv(ws, rows[0], k, v)
    else:
        append(ws, d)


def delete_where(ws, **kv):
    for r in sorted(find(ws, **kv), reverse=True):
        ws.delete_rows(r)


S = wb["sources"]
F = wb["facts"]
O = wb["options"]
SC = wb["schemes"]
C = wb["conditions"]
D = wb["documents"]
SD = wb["scheme_documents"]
A = wb["scheme_apply"]
T = wb["terminology"]
PF = wb["profiles"]


def scheme(sid, **kw):
    base = dict(scheme_id=sid, last_verified=TODAY, verified_by="AI-review")
    base.update(kw)
    upsert(SC, "scheme_id", base)


def cond(sid, cid, group, fact, op, value, quote, locator, source, changeable="no", how_en=None, how_ml=None):
    delete_where(C, scheme_id=sid, condition_id=cid)
    append(C, dict(scheme_id=sid, condition_id=cid, group=group, fact_id=fact, op=op, value=value, changeable=changeable,
                   how_to_en=how_en, how_to_ml=how_ml, source_id=source, quote=quote, locator=locator, reviewed_by=PR))


def docs_for(sid, items):
    delete_where(SD, scheme_id=sid)
    for did, source, quote in items:
        append(SD, dict(scheme_id=sid, doc_id=did, source_id=source, quote=quote))


def apply_for(sid, rows):
    delete_where(A, scheme_id=sid)
    for r in rows:
        append(A, dict(scheme_id=sid, **r))


def doc(did, en, ml, source):
    upsert(D, "doc_id", dict(doc_id=did, name_en=en, name_ml=ml, source_id=source))


def fact(**d):
    upsert(F, "fact_id", d)


BOARD = "Kerala Small Plantation Workers' Welfare Fund Board"
BOARD_ML = "കേരള ചെറുകിട തോട്ടം തൊഴിലാളി ക്ഷേമനിധി ബോർഡ്"
BASE = "https://plantationworker.kerala.gov.in/index.php/benefits/"
BOARD_NOTE = "Read in a browser on the team's computer; the site times out from outside India. Text copied verbatim from the page article."

pages = [
    ("SRC-PLNT-011", "Marriage Assistance", "marriage-assistance"),
    ("SRC-PLNT-012", "Medical Benefit", "medical-benefit"),
    ("SRC-PLNT-013", "Death Assistance", "death-assistance"),
    ("SRC-PLNT-014", "Superannuation", "pension/superannuation"),
    ("SRC-PLNT-015", "Family Pension", "pension/family-pension"),
    ("SRC-PLNT-016", "Invalid Pension", "pension/invalid-pension"),
    ("SRC-PLNT-017", "Maternity Benefit", "maternity-benefit"),
    ("SRC-PLNT-018", "Educational Benefit", "educational-benefit"),
    ("SRC-PLNT-019", "Forms", "forms"),
]
for sid, title, path in pages:
    upsert(S, "source_id", dict(source_id=sid, tier=2, authority=BOARD, title=f"Benefits: {title}", url=BASE + path,
                               doc_type="board_page", language="en", date_issued=None, accessed_on=TODAY, local_file="sources/plantationworker-benefit-pages-2026-09-23.txt", notes=BOARD_NOTE))
r = find(S, source_id="SRC-PLNT-017")[0]
setv(S, r, "notes", BOARD_NOTE + " The membership sentence is printed incomplete: 'at least year' (number missing). Older than, and possibly replaced by, the Labour Commissionerate scheme in G.O.(P) No.81/2024/LBR.")
upsert(S, "source_id", dict(source_id="SRC-LAB-001", tier=2, authority="Labour Commissionerate, Government of Kerala",
       title="Indian Labour Year Book 2016: Kerala contribution (hosted by the Labour Commissionerate)",
       url="https://lc.kerala.gov.in/images/pdf/Report/Indian-Labour-Year-Book-2016.pdf", doc_type="report", language="en",
       date_issued="2016", accessed_on=TODAY, local_file=None,
       notes="Used only to name the Estate Workers (Distress Relief) Welfare Fund Scheme and its relief amount (page 6). Not a rule source; the scheme stays display only."))
r = find(S, source_id="SRC-PLNT-006")[0]
setv(S, r, "notes", "Online membership registration and scheme applications for the 16 Labour Department welfare boards, with Akshaya access. NOT confirmed for the Small Plantation Workers Board: the portal and both help manuals (AIIS_USER_MANUAL_FOR_AKSHAYA.pdf, AIIS_Help_FIle_for_Board_Users.pdf) do not name it, and on 2026-09-23 the public membership registration board list showed 8 boards without it.")

fact(fact_id="plantation_membership_years", type="number", label_en="Years of continuous Welfare Fund membership", label_ml="തുടർച്ചയായ ക്ഷേമനിധി അംഗത്വ വർഷങ്ങൾ",
     question_en="For how many years has the worker been a continuous member of the Small Plantation Workers Welfare Fund?",
     question_ml="തൊഴിലാളി എത്ര വർഷമായി തുടർച്ചയായി ചെറുകിട തോട്ടം തൊഴിലാളി ക്ഷേമനിധി അംഗമാണ്?",
     help_en="Count from the date of joining. A break in membership may restart the count.",
     help_ml="അംഗമായ തീയതി മുതൽ കണക്കാക്കുക. അംഗത്വം മുടങ്ങിയാൽ കാലയളവ് വീണ്ടും തുടങ്ങിയേക്കാം.",
     unit="years", sensitivity="low", ask_order=14)
r = find(F, fact_id="daughter_marriage")[0]
setv(F, r, "label_en", "Marriage of a woman member or a member's daughter")
setv(F, r, "label_ml", "വനിതാ അംഗത്തിന്റെയോ അംഗത്തിന്റെ മകളുടെയോ വിവാഹം")
setv(F, r, "question_en", "Is there a marriage to apply for: a woman member's own marriage, or the marriage of a member's daughter?")
setv(F, r, "question_ml", "അപേക്ഷിക്കാൻ ഒരു വിവാഹമുണ്ടോ: വനിതാ അംഗത്തിന്റെ സ്വന്തം വിവാഹം, അല്ലെങ്കിൽ അംഗത്തിന്റെ മകളുടെ വിവാഹം?")
fact(fact_id="incapacitated_over_2_years", type="boolean", label_en="Permanently unable to work for over two years", label_ml="രണ്ട് വർഷത്തിലധികമായി സ്ഥിരമായി ജോലി ചെയ്യാനാകാത്തത്",
     question_en="Has the member been continuously and permanently unable to work for more than two years?",
     question_ml="അംഗത്തിന് രണ്ട് വർഷത്തിലധികമായി തുടർച്ചയായും സ്ഥിരമായും ജോലി ചെയ്യാൻ കഴിയാത്ത അവസ്ഥയാണോ?",
     sensitivity="high", ask_order=16)
fact(fact_id="plantation_member_death", type="boolean", label_en="Death of a Small Plantation Welfare Fund member", label_ml="ചെറുകിട തോട്ടം ക്ഷേമനിധി അംഗത്തിന്റെ മരണം",
     question_en="Has a member of the Small Plantation Workers Welfare Fund in your family died?",
     question_ml="നിങ്ങളുടെ കുടുംബത്തിലെ ചെറുകിട തോട്ടം തൊഴിലാളി ക്ഷേമനിധി അംഗം മരിച്ചോ?",
     help_en="Under the Act, family means wife or husband, minor sons and unmarried daughters, and dependent parents and dependent disabled or widowed children.",
     help_ml="നിയമപ്രകാരം കുടുംബം: ഭാര്യ/ഭർത്താവ്, മൈനറായ ആൺമക്കൾ, അവിവാഹിതരായ പെൺമക്കൾ, ആശ്രിതരായ മാതാപിതാക്കൾ, ആശ്രിതരായ ഭിന്നശേഷിയുള്ള അല്ലെങ്കിൽ വിധവയായ മക്കൾ.",
     sensitivity="high", ask_order=40)
fact(fact_id="plantation_deceased_contrib_3y", type="boolean", label_en="Deceased member paid contributions for 3 years", label_ml="മരിച്ച അംഗം 3 വർഷം വിഹിതം അടച്ചു",
     question_en="Had the member who died paid contributions to the Fund for at least 3 years?",
     question_ml="മരിച്ച അംഗം കുറഞ്ഞത് 3 വർഷമെങ്കിലും ക്ഷേമനിധിയിൽ വിഹിതം അടച്ചിരുന്നോ?", sensitivity="medium", ask_order=41)
fact(fact_id="plantation_deceased_worked_15y", type="boolean", label_en="Deceased member worked 15 years in a small plantation", label_ml="മരിച്ച അംഗം 15 വർഷം ചെറുകിട തോട്ടത്തിൽ ജോലി ചെയ്തു",
     question_en="Had the member who died worked in a small plantation for at least 15 years?",
     question_ml="മരിച്ച അംഗം കുറഞ്ഞത് 15 വർഷമെങ്കിലും ചെറുകിട തോട്ടത്തിൽ ജോലി ചെയ്തിരുന്നോ?", sensitivity="medium", ask_order=42)
fact(fact_id="plantation_medical_treatment", type="boolean", label_en="Medical treatment for the member or family", label_ml="അംഗത്തിനോ കുടുംബത്തിനോ ചികിൽസ",
     question_en="Has the member or a family member had medical treatment for which you want help with the bills?",
     question_ml="അംഗത്തിനോ കുടുംബാംഗത്തിനോ ചികിൽസ വേണ്ടിവന്നിട്ട് അതിന്റെ ബില്ലുകൾക്ക് സഹായം വേണോ?",
     help_en="The Board asks for original bills, an inpatient card from a Government hospital, and a medical certificate.",
     help_ml="അസ്സൽ ബില്ലുകൾ, സർക്കാർ ആശുപത്രിയിലെ ഇൻപേഷ്യന്റ് കാർഡ്, മെഡിക്കൽ സർട്ടിഫിക്കറ്റ് എന്നിവ ബോർഡ് ആവശ്യപ്പെടുന്നു.",
     sensitivity="medium", ask_order=43)
fact(fact_id="plantation_child_class10", type="boolean", label_en="Member's child in Class X or higher", label_ml="അംഗത്തിന്റെ കുട്ടി പത്താം ക്ലാസിലോ ഉപരിപഠനത്തിലോ",
     question_en="Is a child of the member studying in Class X or in a higher course?",
     question_ml="അംഗത്തിന്റെ കുട്ടി പത്താം ക്ലാസിലോ അതിനു മുകളിലുള്ള കോഴ്സിലോ പഠിക്കുന്നുണ്ടോ?", sensitivity="low", ask_order=30)

r = find(F, fact_id="small_plantation_worker")[0]
setv(F, r, "help_en", "Act 17 of 2008, s.2(m): a person engaged for doing any work in a small plantation for not less than ninety days in the preceding twelve months, including a self-employed person working in his own small plantation of not more than half a hectare. A small plantation is one of less than five hectares (s.2(l)).")
setv(F, r, "help_ml", "2008-ലെ 17-ാം നിയമം, വകുപ്പ് 2(m): കഴിഞ്ഞ 12 മാസത്തിൽ കുറഞ്ഞത് 90 ദിവസം ചെറുകിട തോട്ടത്തിൽ ജോലി ചെയ്തവർ; അര ഹെക്ടറിൽ കൂടാത്ത സ്വന്തം ചെറുകിട തോട്ടത്തിൽ സ്വയം തൊഴിൽ ചെയ്യുന്നവരും ഉൾപ്പെടും. 5 ഹെക്ടറിൽ താഴെയുള്ളതാണ് ചെറുകിട തോട്ടം (വകുപ്പ് 2(l)).")

HOW_JOIN_EN = "If you worked at least 90 days in the last 12 months in a small plantation (under 5 hectares), or run your own plantation of up to half a hectare, ask the District Executive Officer for your area about joining the Fund."
HOW_JOIN_ML = "കഴിഞ്ഞ 12 മാസത്തിൽ കുറഞ്ഞത് 90 ദിവസം 5 ഹെക്ടറിൽ താഴെയുള്ള തോട്ടത്തിൽ ജോലി ചെയ്തെങ്കിൽ, അല്ലെങ്കിൽ അര ഹെക്ടർ വരെയുള്ള സ്വന്തം തോട്ടമുണ്ടെങ്കിൽ, ക്ഷേമനിധിയിൽ ചേരുന്നതിനെക്കുറിച്ച് നിങ്ങളുടെ പ്രദേശത്തെ ജില്ലാ എക്സിക്യൂട്ടീവ് ഓഫീസറോട് അന്വേഷിക്കുക."
Q_M = '"member" means a member of the Fund;'
for sid in ["PLNT-01", "PLNT-02", "PLNT-03", "PLNT-04", "PLNT-05"]:
    for rr in find(C, scheme_id=sid, fact_id="plantation_member"):
        setv(C, rr, "changeable", "yes")
        setv(C, rr, "how_to_en", HOW_JOIN_EN)
        setv(C, rr, "how_to_ml", HOW_JOIN_ML)

DEO = "plantation_district_office"
DEO_QUOTE = "ജില്ലാ എക്സിക്യൂട്ടീവ് ഓഫീസർ"
AK_NOTE_EN = "The Labour Department's AIIS portal (aiis.lc.kerala.gov.in) takes welfare board applications through Akshaya centres. We could not confirm that this Board is on it, so ask at the Akshaya centre before relying on it."
AK_NOTE_ML = "തൊഴിൽ വകുപ്പിന്റെ AIIS പോർട്ടൽ (aiis.lc.kerala.gov.in) അക്ഷയ കേന്ദ്രങ്ങൾ വഴി ക്ഷേമനിധി ബോർഡ് അപേക്ഷകൾ സ്വീകരിക്കുന്നു. ഈ ബോർഡ് അതിലുണ്ടെന്ന് ഉറപ്പാക്കാൻ കഴിഞ്ഞിട്ടില്ല; അക്ഷയ കേന്ദ്രത്തിൽ ചോദിച്ച് ഉറപ്പാക്കുക."
AK_QUOTE = "Akshaya and Trade Union Login / Sign-up"


def deo_row(form_no, source, quote=DEO_QUOTE):
    return dict(type_id=DEO, mode="in_person",
                note_en=f"Submit Form No.{form_no} to the District Executive Officer for your area. The form's office section is filled by the District Executive Officer.",
                note_ml=f"ഫാറം നമ്പർ {form_no} നിങ്ങളുടെ പ്രദേശത്തെ ജില്ലാ എക്സിക്യൂട്ടീവ് ഓഫീസർക്ക് സമർപ്പിക്കുക. ഫാറത്തിലെ ഓഫീസ് ഭാഗം ജില്ലാ എക്സിക്യൂട്ടീവ് ഓഫീസറാണ് പൂരിപ്പിക്കുന്നത്.",
                source_id=source, quote=quote)


def deo_unconfirmed(form_no, source, quote):
    return dict(type_id=DEO, mode="in_person",
                note_en=f"Apply in Form No.{form_no}. The Board page does not say where to submit it; the District Executive Officer for your area handles the Fund's members.",
                note_ml=f"ഫാറം നമ്പർ {form_no}-ൽ അപേക്ഷിക്കുക. എവിടെ സമർപ്പിക്കണമെന്ന് ബോർഡ് പേജിൽ പറയുന്നില്ല; നിങ്ങളുടെ പ്രദേശത്തെ ജില്ലാ എക്സിക്യൂട്ടീവ് ഓഫീസറാണ് ക്ഷേമനിധി അംഗങ്ങളുടെ കാര്യങ്ങൾ കൈകാര്യം ചെയ്യുന്നത്.",
                source_id=source, quote=quote)


AK = dict(type_id="akshaya", mode="online", note_en=AK_NOTE_EN, note_ml=AK_NOTE_ML, source_id="SRC-PLNT-006", quote=AK_QUOTE)
CEO = dict(type_id="spwwfb_office", mode="in_person", note_en="The Chief Executive Officer's office in Kottayam administers the Fund.",
           note_ml="കോട്ടയത്തെ ചീഫ് എക്സിക്യൂട്ടീവ് ഓഫീസറുടെ കാര്യാലയമാണ് ക്ഷേമനിധി ഭരിക്കുന്നത്.", source_id="SRC-PLNT-007",
           quote="Office of the Chief Executive Officer, Azad Lane, Kottayam - 686 001")

for did, en, ml, src in [
    ("spwwfb_form_15", "Form 15: invalid pension application", "ഫാറം 15: അശക്തി പെൻഷനുള്ള അപേക്ഷ", "SRC-PLNT-016"),
    ("spwwfb_form_16", "Form 16: family pension application", "ഫാറം 16: കുടുംബ പെൻഷനുള്ള അപേക്ഷ", "SRC-PLNT-015"),
    ("spwwfb_form_17", "Form 17: maternity benefit application", "ഫാറം 17: പ്രസവാനുകൂല്യത്തിനുള്ള അപേക്ഷ", "SRC-PLNT-017"),
    ("spwwfb_form_19", "Form 19: medical benefit application", "ഫാറം 19: ചികിൽസാ ആനുകൂല്യത്തിനുള്ള അപേക്ഷ", "SRC-PLNT-012"),
    ("spwwfb_education_form", "Educational scholarship application form", "വിദ്യാഭ്യാസ സ്കോളർഷിപ്പിനുള്ള അപേക്ഷാ ഫാറം", "SRC-PLNT-018"),
    ("medical_board_certificate", "Medical certificate issued by a Medical Board", "മെഡിക്കൽ ബോർഡ് നൽകുന്ന മെഡിക്കൽ സർട്ടിഫിക്കറ്റ്", "SRC-PLNT-016"),
    ("original_medical_bills", "Original medical bills", "ചികിൽസയുടെ അസ്സൽ ബില്ലുകൾ", "SRC-PLNT-012"),
    ("govt_inpatient_card", "Inpatient card issued by a Government hospital", "സർക്കാർ ആശുപത്രി നൽകിയ ഇൻപേഷ്യന്റ് കാർഡ്", "SRC-PLNT-012"),
    ("medical_certificate", "Medical certificate", "മെഡിക്കൽ സർട്ടിഫിക്കറ്റ്", "SRC-PLNT-012"),
    ("birth_certificate", "Birth certificate", "ജനന സർട്ടിഫിക്കറ്റ്", "SRC-PLNT-017"),
    ("family_membership_certificate", "Family membership certificate", "കുടുംബാംഗത്വ സർട്ടിഫിക്കറ്റ്", "SRC-PLNT-015"),
]:
    doc(did, en, ml, src)

scheme("PLNT-01", summary_en="Monthly pension for a member who has 5 years of continuous membership, on reaching 60. The minimum pension is ₹1,100 a month (G.O.(MS) No.359/2017/Fin).",
       summary_ml="5 വർഷം തുടർച്ചയായ അംഗത്വമുള്ള അംഗത്തിന് 60 വയസ്സ് തികയുമ്പോൾ പ്രതിമാസ പെൻഷൻ. കുറഞ്ഞ പെൻഷൻ പ്രതിമാസം 1,100 രൂപ (G.O.(MS) No.359/2017/Fin).",
       eligibility_verification="verified", documents_verification="verified", apply_verification="verified",
       notes="Act 17 of 2008 s.3(5)(a) and the Board's Superannuation page (5 years continuous membership, age 60, Form 14). The Board page lists only Form 14; Form 14 itself lists no attachments. Form 14 office section: District Executive Officer. AIIS route not confirmed for this Board.")
cond("PLNT-01", "c3", 3, "plantation_membership_years", "gte", 5, "SUPERANNUATION PENSION: Those who have a continuous membership of 5 years are eligible for the pension on attaining the age of 60.", "Superannuation page", "SRC-PLNT-014")
docs_for("PLNT-01", [("spwwfb_form_14", "SRC-PLNT-014", "Application to be submitted in Form No.14.")])
apply_for("PLNT-01", [deo_row(14, "SRC-PLNT-008"), CEO, AK])

scheme("PLNT-02", summary_en="Pension for a member who has been continuously and permanently unable to work for more than two years.",
       summary_ml="രണ്ട് വർഷത്തിലധികമായി തുടർച്ചയായും സ്ഥിരമായും ജോലി ചെയ്യാൻ കഴിയാത്ത അംഗത്തിന് പെൻഷൻ.",
       eligibility_verification="verified", documents_verification="verified", apply_verification="partial",
       notes="Act 17 of 2008 s.3(5)(a) and the Board's Invalid Pension page. The page sentence is printed as 'Those who have a continuously and permanently incapacitated and unable to do work beyond two years are eligible.' We do not have Form 15, so where to submit it is not confirmed.")
cond("PLNT-02", "c3", 3, "incapacitated_over_2_years", "eq", "yes", "Those who have a continuously and permanently incapacitated and unable to do work beyond two years are eligible.", "Invalid Pension page", "SRC-PLNT-016")
docs_for("PLNT-02", [("spwwfb_form_15", "SRC-PLNT-016", "Application to be submitted in Form No.15."),
                     ("medical_board_certificate", "SRC-PLNT-016", "supported by Medical Certificate issued by a Medical Board.")])
apply_for("PLNT-02", [deo_unconfirmed(15, "SRC-PLNT-016", "Application to be submitted in Form No.15."), CEO, AK])

for sid in ["PLNT-03"]:
    r = find(SC, scheme_id=sid)[0]
    setv(SC, r, "notes", "Act 17 of 2008 s.3(5)(c) only. The Board website has no separate page for permanent-disability assistance (its pension pages cover superannuation, family and invalid pension), so the amount, documents and form are not published.")
    apply_for(sid, [dict(type_id=DEO, mode="in_person", note_en="Ask the District Executive Officer for your area. The Board has not published the application procedure for this benefit.",
                         note_ml="നിങ്ങളുടെ പ്രദേശത്തെ ജില്ലാ എക്സിക്യൂട്ടീവ് ഓഫീസറോട് അന്വേഷിക്കുക. ഈ ആനുകൂല്യത്തിന്റെ അപേക്ഷാ നടപടിക്രമം ബോർഡ് പ്രസിദ്ധീകരിച്ചിട്ടില്ല.",
                         source_id="SRC-PLNT-007", quote="Office of the District Executive Officer"), CEO, AK])

scheme("PLNT-04", name_en="Marriage Assistance (Small Plantation Workers' Welfare Fund)", name_ml="വിവാഹ ധനസഹായം (ചെറുകിട തോട്ടം തൊഴിലാളി ക്ഷേമനിധി)",
       summary_en="₹3,000 for the marriage of a woman member, or of a member's daughter, after at least 3 years of membership.",
       summary_ml="കുറഞ്ഞത് 3 വർഷം അംഗത്വമുള്ള വനിതാ അംഗത്തിന്റെയോ അംഗത്തിന്റെ മകളുടെയോ വിവാഹത്തിന് 3,000 രൂപ.",
       eligibility_verification="verified", documents_verification="verified", apply_verification="verified",
       notes="Act 17 of 2008 s.3(5)(d) and the Board's Marriage Assistance page (female members and members' daughters, Rs 3,000, 3 years, Form 18, marriage certificate). Form 18 office section: District Executive Officer. Form 18 also asks whether the benefit was received or applied for before; no limit on repeat claims is published, so it is not encoded.")
delete_where(C, scheme_id="PLNT-04", condition_id="c2")
cond("PLNT-04", "c2b", 2, "daughter_marriage", "eq", "yes", "This benefit is extended to the female workers who are members of the scheme and to the daughters of the members of the scheme.", "Marriage Assistance page", "SRC-PLNT-011")
cond("PLNT-04", "c3", 3, "plantation_membership_years", "gte", 3, "The workers should have been members of the scheme for at least 3 years.", "Marriage Assistance page", "SRC-PLNT-011")
docs_for("PLNT-04", [("spwwfb_form_18", "SRC-PLNT-011", "Application to be submitted in Form No.18."),
                     ("marriage_certificate", "SRC-PLNT-011", "supported by Marriage Certificate.")])
apply_for("PLNT-04", [deo_row(18, "SRC-PLNT-009"), CEO, AK])

r = find(SC, scheme_id="PLNT-05")[0]
setv(SC, r, "notes", "G.O.(P) No.81/2024/LBR (23 Nov 2024) lists the Small Plantation Workers' Welfare Fund Board (item 10) among the boards whose female members get the Maternity Benefit Scheme, implemented through a web-based Labour Commissionerate application. The amount, number of deliveries and contribution period are in the 'extant Scheme guidelines', not collected. CONFLICT: the Board's own Maternity Benefit page (older) says Rs 15,000 for two deliveries, Rs 1,000 for miscarriage, membership 'at least year' (number missing), Form 17 with birth and medical certificates. The Gazette is newer, so its online route is used; the Board page figures are shown here only as notes, not as rules.")
setv(SC, r, "documents_verification", "partial")
docs_for("PLNT-05", [("aadhaar_or_enrolment", "SRC-PLNT-002", "Any individual eligible for receiving the benefits under the above Scheme shall hereby be required to furnish proof of possession of the Aadhaar number or undergo Aadhaar authentication"),
                     ("birth_certificate", "SRC-PLNT-017", "supported by Birth Certificate and Medical Certificate."),
                     ("medical_certificate", "SRC-PLNT-017", "supported by Birth Certificate and Medical Certificate.")])

scheme("PLNT-06", name_en="Death Assistance to a Member's Dependants (Small Plantation Workers)", name_ml="മരണാനന്തര ധനസഹായം — ആശ്രിതർക്ക് (ചെറുകിട തോട്ടം തൊഴിലാളി)",
       summary_en="₹10,000 to the dependants of a member who has died.", summary_ml="മരിച്ച അംഗത്തിന്റെ ആശ്രിതർക്ക് 10,000 രൂപ.",
       eligibility_verification="verified", documents_verification="verified", apply_verification="verified",
       notes="Board's Death Assistance page and Form 20 (Scheme 2009). No contribution period and no order of priority among heirs is published, so none is encoded. Form 20 asks for the consent of other heirs when more than one is entitled, and has a blank line for certificates attached; the Board page names only Form 20. Form 20 office section: District Executive Officer.")
cond("PLNT-06", "g0", 0, "livelihood", "includes", "plantation_worker", "The dependants of the deceased member of the scheme are eligible to get the benefit of Rs.10000.", "Death Assistance page", "SRC-PLNT-013")
cond("PLNT-06", "c1", 1, "plantation_member_death", "eq", "yes", "The dependants of the deceased member of the scheme are eligible to get the benefit of Rs.10000.", "Death Assistance page", "SRC-PLNT-013")
docs_for("PLNT-06", [("spwwfb_form_20", "SRC-PLNT-013", "Application to be submitted in Form No.20.")])
apply_for("PLNT-06", [deo_row(20, "SRC-PLNT-010"), CEO])

scheme("PLNT-07", summary_en="Family pension at 50% of the member's pension, paid to a family member after the member's death, if the member paid contributions for at least 3 years and worked in a small plantation for at least 15 years.",
       summary_ml="അംഗം മരിച്ചാൽ കുടുംബാംഗത്തിന് അംഗത്തിന്റെ പെൻഷന്റെ 50% കുടുംബ പെൻഷൻ; അംഗം കുറഞ്ഞത് 3 വർഷം വിഹിതം അടയ്ക്കുകയും 15 വർഷം ചെറുകിട തോട്ടത്തിൽ ജോലി ചെയ്യുകയും ചെയ്തിരിക്കണം.",
       eligibility_verification="verified", documents_verification="verified", apply_verification="partial",
       notes="Act 17 of 2008 s.3(5)(b) and the Board's Family Pension page. The Board page asks for an income certificate but publishes no income limit, so none is encoded. We do not have Form 16, so where to submit it is not confirmed.")
cond("PLNT-07", "g0", 0, "livelihood", "includes", "plantation_worker", "for payment of family pension;", "Section 3(5)(b)", ACT)
cond("PLNT-07", "c1", 1, "plantation_member_death", "eq", "yes", "A member of the family is eligible for the family pension on the demise of the member of the Fund @ 50% of the pension amount.", "Family Pension page", "SRC-PLNT-015")
cond("PLNT-07", "c2", 2, "plantation_deceased_contrib_3y", "eq", "yes", "The member should have remitted contribution for a minimum period of 3 years and should have been employed in a small plantation for not less than 15 years.", "Family Pension page", "SRC-PLNT-015")
cond("PLNT-07", "c3", 3, "plantation_deceased_worked_15y", "eq", "yes", "The member should have remitted contribution for a minimum period of 3 years and should have been employed in a small plantation for not less than 15 years.", "Family Pension page", "SRC-PLNT-015")
docs_for("PLNT-07", [("spwwfb_form_16", "SRC-PLNT-015", "Application to be submitted in Form No.16."),
                     ("death_certificate", "SRC-PLNT-015", "supported by Death Certificate, Family Membership Certificate, and Income Certificate."),
                     ("family_membership_certificate", "SRC-PLNT-015", "supported by Death Certificate, Family Membership Certificate, and Income Certificate."),
                     ("income_certificate", "SRC-PLNT-015", "supported by Death Certificate, Family Membership Certificate, and Income Certificate.")])
apply_for("PLNT-07", [deo_unconfirmed(16, "SRC-PLNT-015", "Application to be submitted in Form No.16."), CEO])

scheme("PLNT-08", name_en="Educational Benefit for Members' Children (Small Plantation Workers)", name_ml="അംഗങ്ങളുടെ മക്കൾക്ക് വിദ്യാഭ്യാസ ആനുകൂല്യം (ചെറുകിട തോട്ടം തൊഴിലാളി)",
       authority_en=BOARD, authority_ml=BOARD_ML, category="plantation",
       summary_en="Educational benefit for the children of members with at least 2 years of continuous membership, for courses from Class X onwards. The amount is not published.",
       summary_ml="കുറഞ്ഞത് 2 വർഷം തുടർച്ചയായ അംഗത്വമുള്ളവരുടെ മക്കൾക്ക് പത്താം ക്ലാസ് മുതലുള്ള കോഴ്സുകൾക്ക് വിദ്യാഭ്യാസ ആനുകൂല്യം. തുക പ്രസിദ്ധീകരിച്ചിട്ടില്ല.",
       eligibility_verification="verified", documents_verification="partial", apply_verification="partial",
       notes="Act 17 of 2008 s.3(5)(f) and (h), and the Board's Educational Benefit page. The page gives no amount, no course list and no supporting documents; only an unnumbered application form.")
cond("PLNT-08", "g0", 0, "livelihood", "includes", "plantation_worker", "to provide for the necessities in relation to the education of the Children of Plantation workers;", "Section 3(5)(h)", ACT)
cond("PLNT-08", "c1", 1, "plantation_member", "eq", "yes", Q_M, "Section 2(g)", ACT, changeable="yes", how_en=HOW_JOIN_EN, how_ml=HOW_JOIN_ML)
cond("PLNT-08", "c2", 2, "plantation_membership_years", "gte", 2, "Those workers who have a continuous membership for at least 2 years can avail the benefit for their children for various courses from Class X onwards.", "Educational Benefit page", "SRC-PLNT-018")
cond("PLNT-08", "c3", 3, "plantation_child_class10", "eq", "yes", "Those workers who have a continuous membership for at least 2 years can avail the benefit for their children for various courses from Class X onwards.", "Educational Benefit page", "SRC-PLNT-018")
docs_for("PLNT-08", [("spwwfb_education_form", "SRC-PLNT-019", "Application For Educational Scholarship")])
apply_for("PLNT-08", [dict(type_id=DEO, mode="in_person", note_en="Ask the District Executive Officer for your area for the scholarship form. The Board page does not say where to submit it.",
                           note_ml="സ്കോളർഷിപ്പ് ഫാറത്തിനായി നിങ്ങളുടെ പ്രദേശത്തെ ജില്ലാ എക്സിക്യൂട്ടീവ് ഓഫീസറെ സമീപിക്കുക. എവിടെ സമർപ്പിക്കണമെന്ന് ബോർഡ് പേജിൽ പറയുന്നില്ല.",
                           source_id="SRC-PLNT-007", quote="Office of the District Executive Officer"), CEO])

scheme("PLNT-09", name_en="Medical Benefit (Small Plantation Workers' Welfare Fund)", name_ml="ചികിൽസാ ആനുകൂല്യം (ചെറുകിട തോട്ടം തൊഴിലാളി ക്ഷേമനിധി)",
       authority_en=BOARD, authority_ml=BOARD_ML, category="plantation",
       summary_en="Help with medical bills for the member and family, after at least 3 years of membership. Limited to ₹10,000 in total for the whole period.",
       summary_ml="കുറഞ്ഞത് 3 വർഷം അംഗത്വമുള്ള അംഗത്തിനും കുടുംബത്തിനും ചികിൽസാ ബില്ലുകൾക്ക് സഹായം. ആകെ കാലയളവിൽ പരമാവധി 10,000 രൂപ.",
       eligibility_verification="verified", documents_verification="verified", apply_verification="partial",
       notes="Act 17 of 2008 s.3(5)(f) and the Board's Medical Benefit page. 'Family' follows the Act's definition in s.2(e). The Rs 10,000 lifetime limit is shown in the summary, not asked. We do not have Form 19, so where to submit it is not confirmed.")
cond("PLNT-09", "g0", 0, "livelihood", "includes", "plantation_worker", "for providing facilities for treatment and education to the member and his family members;", "Section 3(5)(f)", ACT)
cond("PLNT-09", "c1", 1, "plantation_member", "eq", "yes", Q_M, "Section 2(g)", ACT, changeable="yes", how_en=HOW_JOIN_EN, how_ml=HOW_JOIN_ML)
cond("PLNT-09", "c2", 2, "plantation_membership_years", "gte", 3, "Those workers who have been a member of the scheme for at least 3 years are entitled to the benefit.", "Medical Benefit page", "SRC-PLNT-012")
cond("PLNT-09", "c3", 3, "plantation_medical_treatment", "eq", "yes", "The worker and his or her family members are entitled to the benefit.", "Medical Benefit page", "SRC-PLNT-012")
Q_MED = "Application to be submitted in Form No.19 supported by the original medical bills, Inpatient Card issued from a Government Hospital and Medical Certificate."
docs_for("PLNT-09", [("spwwfb_form_19", "SRC-PLNT-012", Q_MED), ("original_medical_bills", "SRC-PLNT-012", Q_MED),
                     ("govt_inpatient_card", "SRC-PLNT-012", Q_MED), ("medical_certificate", "SRC-PLNT-012", Q_MED)])
apply_for("PLNT-09", [deo_unconfirmed(19, "SRC-PLNT-012", "Application to be submitted in Form No.19"), CEO, AK])

scheme("PLNT-10", name_en="Estate Workers (Distress Relief) Welfare Fund Scheme, 2007", name_ml="എസ്റ്റേറ്റ് തൊഴിലാളി (ദുരിതാശ്വാസ) ക്ഷേമനിധി പദ്ധതി, 2007",
       authority_en="Labour Commissionerate, Government of Kerala", authority_ml="തൊഴിൽ കമ്മീഷണറേറ്റ്, കേരള സർക്കാർ", category="plantation",
       summary_en="Distress relief of ₹25,000 for each eligible estate worker. Who is eligible is not published in the sources we have.",
       summary_ml="അർഹരായ ഓരോ എസ്റ്റേറ്റ് തൊഴിലാളിക്കും 25,000 രൂപ ദുരിതാശ്വാസം. അർഹത ആർക്കെന്ന് ഞങ്ങളുടെ ഉറവിടങ്ങളിൽ ലഭ്യമല്ല.",
       eligibility_verification="informational", documents_verification="partial", apply_verification="partial",
       notes="Named on the Labour Commissionerate Major Functions page ('A Scheme for giving distress relief to the estate workers') and in the Indian Labour Year Book 2016 contribution hosted by the Commissionerate, p.6. Display only until the scheme notification is found.")
docs_for("PLNT-10", [])
apply_for("PLNT-10", [dict(type_id="labour_commissionerate", mode="in_person", note_en="Ask the Labour Commissionerate or the District Labour Office.",
                           note_ml="തൊഴിൽ കമ്മീഷണറേറ്റിലോ ജില്ലാ ലേബർ ഓഫീസിലോ അന്വേഷിക്കുക.", source_id="SRC-PLNT-004",
                           quote="A Scheme for giving distress relief to the estate workers")])

exec(open(os.path.join(os.path.dirname(__file__), "kfwfb-pages-24-29.py"), encoding="utf-8").read())
exec(open(os.path.join(os.path.dirname(__file__), "profiles-board-pages.py"), encoding="utf-8").read())

wb.save(P)
print("saved")
