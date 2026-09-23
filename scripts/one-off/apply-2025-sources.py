import os
import openpyxl

P = "dataset/anavandi-dataset.xlsx"
wb = openpyxl.load_workbook(P)
TODAY = "2026-09-23"
PR = "PENDING_SECOND_REVIEW"
K = "SRC-KFWFB-001"


def col(ws, name):
    return [c.value for c in ws[1]].index(name) + 1


def ensure_col(ws, name):
    h = [c.value for c in ws[1]]
    if name not in h:
        ws.cell(1, len([x for x in h if x]) + 1).value = name


def find(ws, **kv):
    return [r for r in range(2, ws.max_row + 1) if all(ws.cell(r, col(ws, k)).value == v for k, v in kv.items())]


def append(ws, d):
    r = ws.max_row + 1
    while r > 2 and all(ws.cell(r - 1, c).value is None for c in range(1, ws.max_column + 1)):
        r -= 1
    for k, v in d.items():
        ws.cell(r, col(ws, k)).value = v


def setv(ws, r, k, v):
    ws.cell(r, col(ws, k)).value = v


def get(ws, r, k):
    return ws.cell(r, col(ws, k)).value


def upsert(ws, key, d):
    rows = find(ws, **{key: d[key]})
    if rows:
        for k, v in d.items():
            setv(ws, rows[0], k, v)
    else:
        append(ws, d)


def delete_where(ws, **kv):
    for r in sorted(find(ws, **kv), reverse=True):
        ws.delete_rows(r)


S = wb["sources"]
F = wb["facts"]
O = wb["options"]
SC = wb["schemes"]
C = wb["conditions"]
D = wb["documents"]
SD = wb["scheme_documents"]
LT = wb["location_types"]
L = wb["locations"]
A = wb["scheme_apply"]
T = wb["terminology"]
PF = wb["profiles"]
for c in ["serves_districts", "jurisdiction"]:
    ensure_col(L, c)

# ---------------------------------------------------------------- sources
src = [
    dict(source_id=K, tier=1, authority="Kerala Fishermen's Welfare Fund Board (published by the Fisheries Department)",
         title="Welfare schemes guideline, approved 2025 (68 pages)",
         url="https://fisheries.kerala.gov.in/sites/default/files/inline-files/Scheme%20approved%202025%20KFWFB_0.pdf",
         doc_type="scheme_guideline", language="ml", date_issued="2025", accessed_on=TODAY,
         local_file="sources/kfwfb-welfare-schemes-guideline-2025.pdf",
         notes="Legacy Malayalam font: the text layer is garbled, so quotes were transcribed from the page images (pp.11-31). Newer than the department web pages; preferred where they conflict."),
    dict(source_id="SRC-FISH-012", tier=2, authority="Fisheries Department, Government of Kerala",
         title="Saving cum Relief Scheme application form (Annexure 1, blank-year form)", url="",
         doc_type="other", language="ml", accessed_on=TODAY, local_file="sources/fisheries-saving-cum-relief-application.pdf",
         notes="Collected by hand. 2A.pdf and 2B.pdf were byte-identical; one copy kept. Old blank-year form: its declarations (no mechanised boat, no regular-income job in the family) are NOT encoded as rules until a current guideline confirms them."),
    dict(source_id="SRC-FISH-013", tier=2, authority="Fisheries Department, Government of Kerala",
         title="Saving cum Relief Scheme (inland) credit card", url="", doc_type="other", language="ml", accessed_on=TODAY,
         local_file="sources/fisheries-saving-cum-relief-credit-card.pdf",
         notes="Contribution record card, not an eligibility document. Kept for reference only."),
    dict(source_id="SRC-FISH-014", tier=2, authority="Fisheries Department, Government of Kerala",
         title="Application for compensation to fishermen losing fishing equipments due to calamities",
         url="https://fisheries.kerala.gov.in/sites/default/files/inline-files/appforms/fisheriesdept/Application%20for%20compensation%20to%20fishermen%20losing%20%20fishing%20equipments%20due%20to%20calamities.pdf",
         doc_type="other", language="ml", accessed_on=TODAY, local_file="sources/fisheries-calamity-equipment-loss-application.pdf",
         notes="Scanned form. Gives application fields only; no governing rule found yet."),
    dict(source_id="SRC-PLNT-007", tier=2, authority="Kerala Small Plantation Workers' Welfare Fund Board",
         title="Office list with Inspector of Plantation jurisdiction", url="", doc_type="board_page", language="en",
         accessed_on=TODAY, local_file="sources/spwwfb-offices-jurisdiction.pdf",
         notes="Collected by hand. Jurisdiction is by district and taluk; Vandanmedu is printed as 'Vandanmedu Taluk' and kept as printed."),
    dict(source_id="SRC-PLNT-008", tier=1, authority="Kerala Small Plantation Workers' Welfare Fund Board",
         title="Kerala Small Plantation Workers' Welfare Fund Scheme 2009, Form 14: superannuation pension application (Scheme clause 38(1))",
         url="", doc_type="scheme_guideline", language="ml", accessed_on=TODAY, local_file="sources/spwwfb-form-14-superannuation-pension.pdf",
         notes="Form only; clause 38 itself not collected."),
    dict(source_id="SRC-PLNT-009", tier=1, authority="Kerala Small Plantation Workers' Welfare Fund Board",
         title="Kerala Small Plantation Workers' Welfare Fund Scheme 2009, Form 18: marriage assistance application (Scheme clause 45(1))",
         url="", doc_type="scheme_guideline", language="ml", accessed_on=TODAY, local_file="sources/spwwfb-form-18-marriage-assistance.pdf",
         notes="Form only; clause 45 itself not collected."),
    dict(source_id="SRC-PLNT-010", tier=1, authority="Kerala Small Plantation Workers' Welfare Fund Board",
         title="Kerala Small Plantation Workers' Welfare Fund Scheme 2009, Form 20: assistance to dependants after a member's death (Scheme clause 45(1))",
         url="", doc_type="scheme_guideline", language="ml", accessed_on=TODAY, local_file="sources/spwwfb-form-20-death-assistance.pdf",
         notes="Form only; eligibility and amount are in clause 45, not collected."),
    dict(source_id="SRC-MFED-001", tier=2, authority="Matsyafed (Kerala State Co-operative Federation for Fisheries Development)",
         title="Matsyafed contact list: district offices", url="", doc_type="other", language="en", accessed_on=TODAY,
         local_file="sources/matsyafed-contacts.pdf",
         notes="District offices for 10 districts only (no Pathanamthitta, Idukki, Palakkad or Wayanad). Thiruvananthapuram landline printed incomplete ('0471-'); mobile used."),
]
for d in src:
    upsert(S, "source_id", d)
r = find(S, source_id="SRC-PLNT-001")[0]
setv(S, r, "title", "The Kerala Small Plantation Workers' Welfare Fund Act, 2008 (Act 17 of 2008)")
setv(S, r, "local_file", "sources/kerala-small-plantation-workers-welfare-fund-act-2008.pdf")
setv(S, r, "notes", "Tier 1. The Act PDF itself reads 'Act 17 of 2008'; the India Code listing that shows 18 is treated as a listing error.")
r = find(S, source_id="SRC-PLNT-002")[0]
setv(S, r, "title", "G.O.(P) No.81/2024/LBR dated 23 November 2024, Kerala Gazette Extraordinary 26 November 2024 (Maternity Benefit Scheme, Aadhaar)")
setv(S, r, "date_issued", "2024-11-23")
setv(S, r, "local_file", "sources/gazette-go-p-81-2024-lbr-maternity-aadhaar.pdf")
setv(S, r, "notes", "Tier 1. Lists the Kerala Small Plantation Workers' Welfare Fund Board as item 10 of the 17 boards covered. It is an Aadhaar notification: it does not give the benefit amount or other eligibility rules.")
r = find(S, source_id="SRC-FISH-008")[0]
setv(S, r, "local_file", "sources/fisheries-district-officers.pdf")
setv(S, r, "notes", "Official directory of Fisheries Department district officers for all 14 districts (phones only). Wayanad landline printed malformed (0493-60293214): mobile used. Palakkad email printed as ddfpkd@gmai.com: not used.")
r = find(S, source_id="SRC-FISH-011")[0]
setv(S, r, "local_file", "sources/fisheries-educational-assistance-application.pdf")

# ---------------------------------------------------------------- facts and options
facts = [
    dict(fact_id="kfwfb_contributions_paid", type="boolean", label_en="Welfare Fund contributions paid", label_ml="ക്ഷേമനിധി വിഹിതം അടച്ചു",
         question_en="Are the Fishermen Welfare Fund Board contributions fully paid, with no arrears?",
         question_ml="മത്സ്യത്തൊഴിലാളി ക്ഷേമനിധി വിഹിതം കുടിശ്ശികയില്ലാതെ പൂർണ്ണമായി അടച്ചിട്ടുണ്ടോ?", sensitivity="low", ask_order=20),
    dict(fact_id="kfwfb_membership_years", type="number", label_en="Years since first Welfare Fund contribution", label_ml="ആദ്യ ക്ഷേമനിധി വിഹിതം അടച്ചിട്ട് എത്ര വർഷം",
         question_en="How many years ago was the first Welfare Fund Board contribution paid?",
         question_ml="ക്ഷേമനിധി ബോർഡിൽ ആദ്യമായി വിഹിതം അടച്ചിട്ട് എത്ര വർഷമായി?", unit="years", sensitivity="low", ask_order=21),
    dict(fact_id="fisher_work_10_years", type="boolean", label_en="Worked 10 years as a fisherman in Kerala", label_ml="കേരളത്തിൽ 10 വർഷം മത്സ്യത്തൊഴിലാളി",
         question_en="Has the person worked and earned a living as a fisherman in Kerala for at least 10 years?",
         question_ml="കേരളത്തിൽ കുറഞ്ഞത് 10 വർഷമെങ്കിലും മത്സ്യത്തൊഴിലാളിയായി ജോലി ചെയ്ത് ഉപജീവനം കഴിച്ചിട്ടുണ്ടോ?", sensitivity="low", ask_order=22),
    dict(fact_id="retired_from_fishing", type="boolean", label_en="Retired from fishing work", label_ml="തൊഴിലിൽ നിന്ന് വിരമിച്ചു",
         question_en="Has the person stopped (retired from) fishing work?", question_ml="മത്സ്യബന്ധന തൊഴിലിൽ നിന്ന് വിരമിച്ചോ?", sensitivity="low", ask_order=23),
    dict(fact_id="annual_family_income", type="number", label_en="Yearly family income", label_ml="കുടുംബ വാർഷിക വരുമാനം",
         question_en="What is the family's total income in a year?", question_ml="കുടുംബത്തിന്റെ ഒരു വർഷത്തെ ആകെ വരുമാനം എത്ര?",
         help_en="As shown on the ration card or income certificate.", help_ml="റേഷൻ കാർഡിലോ വരുമാന സർട്ടിഫിക്കറ്റിലോ ഉള്ളത്.",
         unit="INR", sensitivity="medium", ask_order=24),
    dict(fact_id="kfwfb_family_status", type="multi", label_en="Fisherman in the family", label_ml="കുടുംബത്തിലെ മത്സ്യത്തൊഴിലാളി",
         question_en="Which of these describe your family? Choose all that apply.", question_ml="ഇവയിൽ ഏതെല്ലാമാണ് നിങ്ങളുടെ കുടുംബത്തിന് ബാധകം? ബാധകമായതെല്ലാം തിരഞ്ഞെടുക്കുക.",
         sensitivity="medium", ask_order=25),
    dict(fact_id="prior_marriage_assistance", type="boolean", label_en="Received marriage assistance before", label_ml="മുമ്പ് വിവാഹ ധനസഹായം ലഭിച്ചു",
         question_en="Has the family received this Welfare Fund marriage assistance before?", question_ml="ഈ ക്ഷേമനിധി വിവാഹ ധനസഹായം മുമ്പ് ലഭിച്ചിട്ടുണ്ടോ?", sensitivity="low", ask_order=26),
    dict(fact_id="daughter_marriage_60_days", type="boolean", label_en="Daughter's marriage (apply within 60 days)", label_ml="മകളുടെ വിവാഹം (60 ദിവസത്തിനകം)",
         question_en="Is a daughter getting married soon, or did she marry within the last 60 days?", question_ml="മകളുടെ വിവാഹം അടുത്ത് നടക്കാനുണ്ടോ, അല്ലെങ്കിൽ കഴിഞ്ഞ 60 ദിവസത്തിനുള്ളിൽ നടന്നോ?", sensitivity="low", ask_order=27),
    dict(fact_id="bride_18", type="boolean", label_en="Bride 18 on the marriage date", label_ml="വിവാഹ തീയതിക്ക് വധുവിന് 18 വയസ്സ്",
         question_en="Will the bride have completed 18 years on the marriage date?", question_ml="വിവാഹ തീയതിക്ക് വധുവിന് 18 വയസ്സ് പൂർത്തിയാകുമോ?", sensitivity="low", ask_order=28),
    dict(fact_id="member_death", type="boolean", label_en="Death of a Welfare Fund member", label_ml="ക്ഷേമനിധി അംഗത്തിന്റെ മരണം",
         question_en="Has a fisherman in your family who was a Welfare Fund Board member died?", question_ml="നിങ്ങളുടെ കുടുംബത്തിലെ ക്ഷേമനിധി അംഗമായ മത്സ്യത്തൊഴിലാളി മരിച്ചോ?", sensitivity="high", ask_order=40),
    dict(fact_id="death_by_accident", type="boolean", label_en="Death caused by an accident", label_ml="അപകടം മൂലമുള്ള മരണം",
         question_en="Was the death caused by an accident?", question_ml="മരണം അപകടം മൂലമായിരുന്നോ?", sensitivity="high", ask_order=41),
    dict(fact_id="deceased_under_60", type="boolean", label_en="Deceased was under 60", label_ml="മരിച്ചയാൾക്ക് 60 വയസ്സിൽ താഴെ",
         question_en="Was the fisherman below 60 years of age when he died?", question_ml="മരിക്കുമ്പോൾ മത്സ്യത്തൊഴിലാളിക്ക് 60 വയസ്സിൽ താഴെയായിരുന്നോ?", sensitivity="high", ask_order=42),
    dict(fact_id="deceased_active_after_60", type="boolean", label_en="Still working and contributing after 60", label_ml="60-ന് ശേഷവും തൊഴിലും വിഹിതവും",
         question_en="If he was 60 or older: was he still working, on the fishermen list, and paying contributions regularly?", question_ml="60 വയസ്സോ അതിൽ കൂടുതലോ ആയിരുന്നെങ്കിൽ: തൊഴിൽ തുടരുകയും പട്ടികയിൽ പേരുണ്ടായിരിക്കുകയും വിഹിതം കൃത്യമായി അടക്കുകയും ചെയ്തിരുന്നോ?", sensitivity="high", ask_order=43),
    dict(fact_id="death_within_3_months", type="boolean", label_en="Death within the last 3 months", label_ml="കഴിഞ്ഞ 3 മാസത്തിനുള്ളിലെ മരണം",
         question_en="Did the death happen within the last 3 months (90 days)?", question_ml="മരണം കഴിഞ്ഞ 3 മാസത്തിനുള്ളിൽ (90 ദിവസം) ആയിരുന്നോ?", sensitivity="high", ask_order=44),
    dict(fact_id="regular_income_job", type="boolean", label_en="Widow has a regular-income job", label_ml="സ്ഥിര വരുമാനമുള്ള ഉദ്യോഗം",
         question_en="Does the widow have a job with a regular income?", question_ml="വിധവയ്ക്ക് സ്ഥിര വരുമാനമുള്ള ഉദ്യോഗമുണ്ടോ?", sensitivity="medium", ask_order=45),
    dict(fact_id="widow_remarried", type="boolean", label_en="Widow has remarried", label_ml="വിധവ പുനർവിവാഹം ചെയ്തു",
         question_en="Has the widow remarried?", question_ml="വിധവ പുനർവിവാഹം ചെയ്തിട്ടുണ്ടോ?", sensitivity="high", ask_order=46),
    dict(fact_id="dependant_died", type="boolean", label_en="Death of a fisherman's dependant", label_ml="മത്സ്യത്തൊഴിലാളിയുടെ ആശ്രിതന്റെ മരണം",
         question_en="In the last 3 months, has a fisherman's father, mother, wife or husband, minor son or unmarried daughter died?",
         question_ml="കഴിഞ്ഞ 3 മാസത്തിനുള്ളിൽ മത്സ്യത്തൊഴിലാളിയുടെ അച്ഛൻ, അമ്മ, ഭാര്യ/ഭർത്താവ്, മൈനറായ മകൻ, അവിവാഹിതയായ മകൾ എന്നിവരിൽ ആരെങ്കിലും മരിച്ചോ?", sensitivity="high", ask_order=47),
    dict(fact_id="accident_unable_7_days", type="boolean", label_en="Unable to work 7+ days after an accident", label_ml="അപകടം മൂലം 7 ദിവസത്തിലധികം ജോലി ചെയ്യാനായില്ല",
         question_en="Because of an accident, has the fisherman been unable to work for 7 days or more?", question_ml="അപകടം മൂലം മത്സ്യത്തൊഴിലാളിക്ക് 7 ദിവസമോ അതിൽ കൂടുതലോ ജോലി ചെയ്യാൻ കഴിയാതെ വന്നോ?", sensitivity="high", ask_order=48),
    dict(fact_id="accident_within_3_months", type="boolean", label_en="Accident within the last 3 months", label_ml="കഴിഞ്ഞ 3 മാസത്തിനുള്ളിലെ അപകടം",
         question_en="Did the accident happen within the last 3 months?", question_ml="അപകടം കഴിഞ്ഞ 3 മാസത്തിനുള്ളിലായിരുന്നോ?", sensitivity="medium", ask_order=49),
    dict(fact_id="child_sslc_8_aplus", type="boolean", label_en="Child with A+ in 8+ SSLC subjects", label_ml="എസ്.എസ്.എൽ.സി-യിൽ 8+ വിഷയങ്ങളിൽ എ+",
         question_en="Did a child in the family get A+ in at least 8 subjects in this year's SSLC exam?", question_ml="ഈ വർഷത്തെ എസ്.എസ്.എൽ.സി പരീക്ഷയിൽ കുടുംബത്തിലെ കുട്ടിക്ക് കുറഞ്ഞത് 8 വിഷയങ്ങളിൽ എ+ ലഭിച്ചോ?", sensitivity="low", ask_order=30),
    dict(fact_id="serious_illness", type="boolean", label_en="Serious illness needing treatment", label_ml="മാരകരോഗ ചികിൽസ ആവശ്യം",
         question_en="Does the person need treatment for heart disease, kidney disease, cancer, a brain tumour, paralysis, or a mental illness that can be cured with treatment?",
         question_ml="ഹൃദ്രോഗം, വൃക്കരോഗം, ക്യാൻസർ, തലച്ചോറിലെ ട്യൂമർ, തളർവാതം, ചികിൽസിച്ച് ഭേദമാക്കാൻ പറ്റുന്ന മാനസികരോഗം എന്നിവയിൽ ഏതിനെങ്കിലും ചികിൽസ ആവശ്യമുണ്ടോ?", sensitivity="high", ask_order=50),
]
for d in facts:
    upsert(F, "fact_id", d)
r = find(F, fact_id="kfwfb_member")[0]
setv(F, r, "question_en", "Is the person a member of the Kerala Fishermen's Welfare Fund Board (Matsyaboard)?")
setv(F, r, "question_ml", "കേരള മത്സ്യത്തൊഴിലാളി ക്ഷേമനിധി ബോർഡിൽ (മത്സ്യബോർഡ്) അംഗമാണോ?")
r = find(F, fact_id="age")[0]
setv(F, r, "question_en", "How old is the person applying?")
setv(F, r, "question_ml", "അപേക്ഷിക്കുന്ന വ്യക്തിയുടെ പ്രായം എത്ര?")
r = find(F, fact_id="plantation_member")[0]
setv(F, r, "help_en", "The Act covers small plantation workers: people who worked at least 90 days in the last 12 months in a plantation of less than 5 hectares (rubber, tea, coffee, cardamom, coco, oil palm or cashew), or self-employed on their own plantation of up to half a hectare.")
setv(F, r, "help_ml", "5 ഹെക്ടറിൽ താഴെയുള്ള തോട്ടത്തിൽ (റബ്ബർ, തേയില, കാപ്പി, ഏലം, കൊക്കോ, എണ്ണപ്പന, കശുമാവ്) കഴിഞ്ഞ 12 മാസത്തിൽ കുറഞ്ഞത് 90 ദിവസം ജോലി ചെയ്തവർ, അല്ലെങ്കിൽ അര ഹെക്ടർ വരെയുള്ള സ്വന്തം തോട്ടത്തിൽ സ്വയം തൊഴിൽ ചെയ്യുന്നവർ.")
for fid in ["fisher_beneficiary_status"]:
    delete_where(F, fact_id=fid)
    delete_where(O, fact_id=fid)
opts = [
    ("member_under_60", "Board member fisherman, under 60", "60 വയസ്സിൽ താഴെയുള്ള ബോർഡ് അംഗമായ മത്സ്യത്തൊഴിലാളി"),
    ("member_over_60", "Board member fisherman, 60 or older", "60 വയസ്സോ അതിൽ കൂടുതലോ ഉള്ള ബോർഡ് അംഗമായ മത്സ്യത്തൊഴിലാളി"),
    ("pensioner", "Receives the Board's fishermen pension", "ബോർഡിന്റെ മത്സ്യത്തൊഴിലാളി പെൻഷൻ വാങ്ങുന്നയാൾ"),
    ("widow_of_fisher", "Widow of a fisherman who was a Board member", "ബോർഡ് അംഗമായിരുന്ന മത്സ്യത്തൊഴിലാളിയുടെ വിധവ"),
    ("widow_of_pensioner", "Widow of a Board pensioner", "ബോർഡ് പെൻഷൻകാരന്റെ വിധവ"),
    ("bride_is_fisher", "The bride herself is a fisherworker", "വധു തന്നെ മത്സ്യത്തൊഴിലാളിയാണ്"),
    ("fisher_parent_died", "The bride's fisherman parent has died", "വധുവിന്റെ മത്സ്യത്തൊഴിലാളിയായ അച്ഛനോ അമ്മയോ മരിച്ചു"),
    ("none", "None of these", "ഇവയൊന്നുമല്ല"),
]
delete_where(O, fact_id="kfwfb_family_status")
for v, en, ml in opts:
    append(O, dict(fact_id="kfwfb_family_status", value=v, label_en=en, label_ml=ml))

# ---------------------------------------------------------------- documents
docs = [
    ("death_certificate", "Death certificate (original)", "മരണ സർട്ടിഫിക്കറ്റ് (അസ്സൽ)"),
    ("fir_copy", "Copy of the police First Information Report (FIR)", "പോലീസ് സ്റ്റേഷനിലെ പ്രഥമ വിവര റിപ്പോർട്ടിന്റെ (എഫ് ഐ ആർ) പകർപ്പ്"),
    ("postmortem_inquest", "Copies of the post-mortem certificate and inquest report", "പോസ്റ്റ്മോർട്ടം സർട്ടിഫിക്കറ്റിന്റെയും ഇൻക്വസ്റ്റ് റിപ്പോർട്ടിന്റെയും പകർപ്പ്"),
    ("legal_heir_certificate", "Legal heirship / relationship certificate", "ലീഗൽ ഹെയർഷിപ്പ് / റിലേഷൻഷിപ്പ് സർട്ടിഫിക്കറ്റ്"),
    ("student_school_certificate", "Student's school certificate (for the children's education benefit)", "വിദ്യാഭ്യാസ ആനുകൂല്യത്തിന് വിദ്യാർത്ഥിയുടെ സ്കൂൾ സർട്ടിഫിക്കറ്റ്"),
    ("funeral_bill", "Bill for the funeral expense", "ശവസംസ്കാര ചെലവ് ലഭ്യമാകുന്നതിന് ആയതിന്റെ ബിൽ"),
    ("disability_certificate", "Hospital case certificate / disability certificate", "ആശുപത്രിയിലെ കേസ് സർട്ടിഫിക്കറ്റ്/ഡിസെബിലിറ്റി സർട്ടിഫിക്കറ്റ്"),
    ("doctor_certificate_form", "Doctor's certificate in the prescribed form", "ഡോക്ടറുടെ സർട്ടിഫിക്കറ്റ് നിശ്ചിത ഫോറത്തിൽ"),
    ("police_report", "Report from the police station", "പോലീസ് സ്റ്റേഷനിൽ നിന്നുള്ള റിപ്പോർട്ട്"),
    ("collector_rdo_certificate", "Certificate from the Collector / RDO (missing at sea)", "കളക്ടർ/ആർ.ഡി.ഒ നൽകുന്ന സർട്ടിഫിക്കറ്റ്"),
    ("hospital_discharge_card", "Hospital discharge card", "ആശുപത്രി ഡിസ്ചാർജ്ജ് കാർഡ്"),
    ("hospital_bill", "Original hospital bill", "ആശുപത്രിയിലെ അസ്സൽ ബിൽ"),
    ("kfwfb_passbook", "Copy of the Matsyaboard passbook, attested by the Fisheries Officer", "ഫിഷറീസ് ഓഫീസർ സാക്ഷ്യപ്പെടുത്തിയ മത്സ്യബോർഡ് പാസ്സ് ബുക്കിന്റെ പകർപ്പ്"),
    ("bank_passbook", "Copy of the bank passbook", "ബാങ്ക് പാസ് ബുക്കിന്റെ പകർപ്പ്"),
    ("bank_passbook_aadhaar_linked", "Copy of the personal bank passbook (Aadhaar-linked account)", "ആധാർ ലിങ്ക് ചെയ്ത വ്യക്തിഗത ബാങ്ക് അക്കൗണ്ട് പാസ് ബുക്കിന്റെ പകർപ്പ്"),
    ("ration_card", "Copy of the ration card", "റേഷൻ കാർഡിന്റെ പകർപ്പ്"),
    ("income_certificate", "Income certificate", "വരുമാന സർട്ടിഫിക്കറ്റ്"),
    ("income_proof_ration_or_certificate", "Ration card or income certificate copy showing income", "വരുമാനം തെളിയിക്കുന്ന റേഷൻ കാർഡ് / വരുമാന സർട്ടിഫിക്കറ്റ് പകർപ്പ്"),
    ("age_proof", "Age proof: school or birth-registrar certificate, passport, driving licence, or medical certificate in the prescribed form", "വയസ്സ് തെളിയിക്കുന്ന രേഖ: സ്കൂൾ / ജനന മരണ രജിസ്ട്രാർ സർട്ടിഫിക്കറ്റ്, പാസ്പോർട്ട്, ഡ്രൈവിംഗ് ലൈസൻസ്, അല്ലെങ്കിൽ നിശ്ചിത ഫോറത്തിലുള്ള മെഡിക്കൽ സർട്ടിഫിക്കറ്റ്"),
    ("bride_age_proof", "Proof of the bride's age", "വധുവിന്റെ വയസ്സ് തെളിയിക്കുന്ന രേഖ"),
    ("marriage_certificate", "Marriage certificate from the local self-government institution", "തദ്ദേശസ്വയംഭരണ സ്ഥാപനങ്ങളിൽ നിന്നും ലഭിക്കുന്ന വിവാഹ സർട്ടിഫിക്കറ്റ്"),
    ("aadhaar_copy", "Copy of the Aadhaar card", "ആധാർ കാർഡിന്റെ പകർപ്പ്"),
    ("doctor_certificate_cause", "Certificate from the treating doctor on the cause of death (if the death was after hospital admission)", "ചികിൽസിച്ച ഡോക്ടറുടെ സർട്ടിഫിക്കറ്റ് (ആശുപത്രിയിൽ പ്രവേശിച്ച ശേഷമാണ് മരണമെങ്കിൽ)"),
    ("bank_details", "Bank account details of the claimant", "അവകാശിയുടെ ബാങ്ക് അക്കൗണ്ട് വിശദാംശങ്ങൾ"),
    ("dependant_death_certificate", "Death certificate of the dependant", "ആശ്രിതന്റെ മരണ സർട്ടിഫിക്കറ്റ്"),
    ("govt_medical_certificate", "Medical certificate in the prescribed form from the treating government doctor", "ചികിൽസിച്ച സർക്കാർ ഡോക്ടറിൽ നിന്നും ലഭിച്ച നിശ്ചിത ഫോറത്തിലുള്ള മെഡിക്കൽ സർട്ടിഫിക്കറ്റ്"),
    ("sslc_certificate", "Copy of the SSLC certificate, attested by the Fisheries Officer", "ഫിഷറീസ് ഓഫീസർ സാക്ഷ്യപ്പെടുത്തിയ എസ്.എസ്.എൽ.സി സർട്ടിഫിക്കറ്റിന്റെ ശരിപ്പകർപ്പ്"),
    ("parent_kfwfb_passbook", "Copy of the parent's Matsyaboard passbook, attested by the Fisheries Officer", "ഫിഷറീസ് ഓഫീസർ സാക്ഷ്യപ്പെടുത്തിയ രക്ഷിതാക്കളുടെ മത്സ്യബോർഡ് പാസ്സ് ബുക്കിന്റെ ശരി പകർപ്പ്"),
    ("photos_2", "Two passport-size photographs of the student", "വിദ്യാർത്ഥിയുടെ പാസ്പോർട്ട് സൈസ് ഫോട്ടോ 2 എണ്ണം"),
    ("govt_hospital_medical_certificate", "Medical certificate from the doctor who treated at a government hospital", "സർക്കാർ ആശുപത്രിയിൽ ചികിൽസിച്ച ഡോക്ടർ നൽകുന്ന മെഡിക്കൽ സർട്ടിഫിക്കറ്റ്"),
    ("referral_document", "Referral document, if referred from a government hospital to a private hospital", "സർക്കാർ ആശുപത്രിയിൽ നിന്നും സ്വകാര്യ ആശുപത്രിയിലേക്ക് റഫർ ചെയ്തിട്ടുണ്ടെങ്കിൽ ആയതിനുള്ള രേഖ"),
    ("treatment_bills", "Treatment bills attested by the treating doctor", "ചികിൽസിച്ച ഡോക്ടർ സാക്ഷ്യപ്പെടുത്തിയ ചികിൽസാ ചിലവിലേക്കുള്ള ബില്ലുകൾ"),
    ("patient_age_proof", "Patient's age proof: school certificate, birth-registrar certificate or baptism certificate", "രോഗിയുടെ വയസ്സ് തെളിയിക്കുന്നതിന് സ്കൂൾ സർട്ടിഫിക്കറ്റ് അഥവാ ജനന മരണ രജിസ്ട്രാറുടെ സർട്ടിഫിക്കറ്റ് ഇല്ലെങ്കിൽ ജ്ഞാനസ്നാന സർട്ടിഫിക്കറ്റ്"),
    ("prescription_copy", "Copy of the treating doctor's prescription", "ചികിൽസിക്കുന്ന ഡോക്ടറുടെ പ്രിസ്ക്രിപ്ഷന്റെ പകർപ്പ്"),
    ("member_death_certificate_attested", "Copy of the fisherman's death certificate, attested by the Fisheries Officer", "ഫിഷറീസ് ഓഫീസർ സാക്ഷ്യപ്പെടുത്തിയ മത്സ്യത്തൊഴിലാളിയുടെ മരണ സർട്ടിഫിക്കറ്റ് പകർപ്പ്"),
    ("passbook_contribution_pages", "Full copy of the Matsyaboard passbook including contribution pages, attested by the Fisheries Officer", "ഫിഷറീസ് ഓഫീസർ സാക്ഷ്യപ്പെടുത്തിയ മത്സ്യബോർഡ് വിഹിതമടച്ച പേജിന്റെയും പാസ്സ് ബുക്കിന്റെയും പൂർണ്ണമായ പകർപ്പ്"),
    ("no_other_pension_proof", "Proof of not receiving any other pension", "മറ്റു പെൻഷൻ ലഭിക്കുന്നില്ല എന്ന് തെളിയിക്കുന്ന രേഖ"),
    ("ration_card_family_page", "Copy of the ration card including the family details page", "കുടുംബ വിവരം അടങ്ങിയ പേജ് സഹിതം റേഷൻ കാർഡിന്റെ പകർപ്പ്"),
    ("not_remarried_certificate", "Certificate from the panchayat / municipality / corporation secretary or ward councillor that she has not remarried and has no regular-income job", "പുനർവിവാഹം നടത്തിയിട്ടില്ലെന്നും സ്ഥിരവരുമാനമുള്ള തൊഴിൽ ഇല്ലെന്നും തെളിയിക്കുന്നതിന് പഞ്ചായത്ത്, മുൻസിപ്പാലിറ്റി / കോർപ്പറേഷൻ സെക്രട്ടറി / വാർഡ് കൗൺസിലർ നൽകുന്ന സാക്ഷ്യപത്രം"),
    ("scr_contribution_receipt", "Receipt number of this year's Welfare Fund contribution (asked in the form)", "നടപ്പു വർഷ ക്ഷേമനിധി വിഹിതം ഒടുക്കിയതിന്റെ രസീത് നമ്പർ"),
    ("calamity_application", "Application for compensation to fishermen losing fishing equipment due to calamities", "പ്രകൃതിക്ഷോഭം മൂലം മത്സ്യബന്ധന ഉപകരണങ്ങൾ നഷ്ടപ്പെട്ടവർക്കുള്ള നഷ്ടപരിഹാര അപേക്ഷ"),
    ("spwwfb_form_14", "Form 14: superannuation pension application (Scheme 2009)", "ഫാറം 14: സൂപ്പറാന്വേഷൻ പെൻഷനുള്ള അപേക്ഷ (പദ്ധതി 2009)"),
    ("spwwfb_form_18", "Form 18: marriage assistance application (Scheme 2009)", "ഫാറം 18: വിവാഹ ധനസഹായത്തിനുള്ള അപേക്ഷ (പദ്ധതി 2009)"),
    ("spwwfb_form_20", "Form 20: application for assistance to dependants after a member's death (Scheme 2009)", "ഫാറം 20: അംഗത്തിന്റെ മരണാനന്തരം ആശ്രിതർക്ക് ധനസഹായത്തിനുള്ള അപേക്ഷ (പദ്ധതി 2009)"),
    ("aadhaar_or_enrolment", "Aadhaar; or, if not yet assigned, the Aadhaar enrolment slip with one other identity document (for example bank or post office passbook with photo, PAN card, passport, ration card, voter ID, MGNREGA card, Kisan photo passbook)", "ആധാർ; ഇല്ലെങ്കിൽ ആധാർ എൻറോൾമെന്റ് സ്ലിപ്പും മറ്റൊരു തിരിച്ചറിയൽ രേഖയും (ഫോട്ടോയുള്ള ബാങ്ക്/പോസ്റ്റ് ഓഫീസ് പാസ്ബുക്ക്, പാൻ കാർഡ്, പാസ്പോർട്ട്, റേഷൻ കാർഡ്, വോട്ടർ ഐഡി, തൊഴിലുറപ്പ് കാർഡ്, കിസാൻ ഫോട്ടോ പാസ്ബുക്ക് തുടങ്ങിയവ)"),
]
issuers = {
    "death_certificate": ("Local body (panchayat / municipality)", "തദ്ദേശ സ്ഥാപനം"),
    "income_certificate": ("Village Office", "വില്ലേജ് ഓഫീസ്"),
    "marriage_certificate": ("Local body (panchayat / municipality)", "തദ്ദേശ സ്ഥാപനം"),
}
for did, en, ml in docs:
    d = dict(doc_id=did, name_en=en, name_ml=ml, source_id=K)
    if did.startswith("spwwfb_form_14"):
        d["source_id"] = "SRC-PLNT-008"
    if did.startswith("spwwfb_form_18"):
        d["source_id"] = "SRC-PLNT-009"
    if did.startswith("spwwfb_form_20"):
        d["source_id"] = "SRC-PLNT-010"
    if did == "calamity_application":
        d["source_id"] = "SRC-FISH-014"
    if did == "scr_contribution_receipt":
        d["source_id"] = "SRC-FISH-012"
    if did == "aadhaar_or_enrolment":
        d["source_id"] = "SRC-PLNT-002"
    upsert(D, "doc_id", d)

# ---------------------------------------------------------------- helpers for schemes
def scheme(sid, **kw):
    base = dict(scheme_id=sid, last_verified=TODAY, verified_by="AI-review")
    base.update(kw)
    upsert(SC, "scheme_id", base)


def cond(sid, cid, group, fact, op, value, quote, locator, source=K, changeable="no", how_en=None, how_ml=None):
    delete_where(C, scheme_id=sid, condition_id=cid)
    append(C, dict(scheme_id=sid, condition_id=cid, group=group, fact_id=fact, op=op, value=value, changeable=changeable,
                   how_to_en=how_en, how_to_ml=how_ml, source_id=source, quote=quote, locator=locator, reviewed_by=PR))


def docs_for(sid, items, source=K, locator=None):
    delete_where(SD, scheme_id=sid)
    for did, quote in items:
        append(SD, dict(scheme_id=sid, doc_id=did, source_id=source, quote=quote))


def apply_for(sid, rows):
    delete_where(A, scheme_id=sid)
    for r in rows:
        append(A, r)


FISH_OFFICE = "fisheries_district_office"
HOW_KFWFB_EN = "Ask at your Fisheries Department district office about joining the Kerala Fishermen's Welfare Fund Board."
HOW_KFWFB_ML = "കേരള മത്സ്യത്തൊഴിലാളി ക്ഷേമനിധി ബോർഡിൽ അംഗമാകുന്നതിനെക്കുറിച്ച് ഫിഷറീസ് വകുപ്പ് ജില്ലാ ഓഫീസിൽ അന്വേഷിക്കുക."
HOW_CONTRIB_EN = "Clear any arrears of Welfare Fund contribution at the Fisheries Office."
HOW_CONTRIB_ML = "ക്ഷേമനിധി വിഹിതത്തിലെ കുടിശ്ശിക ഫിഷറീസ് ഓഫീസിൽ അടച്ചുതീർക്കുക."

# ---------------------------------------------------------------- FISH-01 group accident insurance (2025)
scheme("FISH-01", name_en="Group Accident Insurance Scheme", name_ml="ഗ്രൂപ്പ് അപകട ഇൻഷ്വറൻസ് പദ്ധതി",
       authority_en="Kerala Fishermen's Welfare Fund Board", authority_ml="കേരള മത്സ്യത്തൊഴിലാളി ക്ഷേമനിധി ബോർഡ്",
       summary_en="Accident cover for every Board member, with no premium from the family: ₹10 lakh for accidental death, missing while fishing, or permanent total disability; ₹5 lakh for permanent partial disability; up to ₹25,000 when an accident needs more than 24 hours in hospital; ₹5,000 each for the education of 2 children of a member who died in an accident (up to ₹10,000); ₹2,500 funeral expense.",
       summary_ml="ബോർഡ് അംഗങ്ങൾക്കെല്ലാം അപകട പരിരക്ഷ; ഗുണഭോക്താവിൽ നിന്ന് തുക ഈടാക്കുന്നില്ല: അപകടമരണം, മത്സ്യബന്ധനവേളയിൽ കാണാതാകൽ, സ്ഥിരവും പൂർണ്ണവുമായ അവശത എന്നിവയ്ക്ക് 10 ലക്ഷം രൂപ; സ്ഥിരവും ഭാഗികവുമായ അവശതയ്ക്ക് 5 ലക്ഷം; 24 മണിക്കൂറിലധികം ആശുപത്രി ചികിൽസയ്ക്ക് പരമാവധി 25,000; അപകടമരണമടഞ്ഞ മത്സ്യത്തൊഴിലാളികളുടെ 2 കുട്ടികൾക്ക് 5,000 വീതം (പരമാവധി 10,000) വിദ്യാഭ്യാസ സഹായം; ശവസംസ്കാര ചെലവ് 2,500.",
       eligibility_verification="verified", documents_verification="verified", apply_verification="verified",
       notes="Upgraded to the 2025 KFWFB guideline p.11-12. CONFLICT: the older Fisheries Department page says 'Active Fishermen' are insured with Rs 1 lakh cover; the 2025 guideline says all Board members, Rs 10 lakh. The 2025 guideline is used. Documents depend on the claim type and are named that way.")
delete_where(C, scheme_id="FISH-01", condition_id="c1")
cond("FISH-01", "c1", 1, "kfwfb_member", "eq", "yes", "മത്സ്യബോർഡിൽ അംഗത്വമുള്ള എല്ലാ മത്സ്യത്തൊഴിലാളികളെയും ഈ പദ്ധതി പ്രകാരം ഇൻഷുർ ചെയ്തുവരുന്നു.", "p.11, para 1",
     changeable="yes", how_en=HOW_KFWFB_EN, how_ml=HOW_KFWFB_ML)
cond("FISH-01", "g0", 0, "livelihood", "includes", "fishing", "പരമ്പരാഗത മത്സ്യത്തൊഴിലാളികളുടെ കുടുംബ ഭദ്രത ഉറപ്പുവരുത്തുന്നതിനുവേണ്ടി മത്സ്യബോർഡിന്റെ തുടക്കം മുതൽ നടപ്പാക്കി വരുന്ന പദ്ധതിയാണിത്.", "p.11, para 1")
docs_for("FISH-01", [
    ("death_certificate", "മരണ സർട്ടിഫിക്കറ്റ് അസ്സൽ"),
    ("fir_copy", "പോലീസ് സ്റ്റേഷനിലെ പ്രഥമ വിവര റിപ്പോർട്ടിന്റെ (എഫ് ഐ ആർ) പകർപ്പ്"),
    ("postmortem_inquest", "പോസ്റ്റ്മോർട്ടം സർട്ടിഫിക്കറ്റിന്റെയും ഇൻക്വസ്റ്റ് റിപ്പോർട്ടിന്റെയും പകർപ്പ്"),
    ("legal_heir_certificate", "ലീഗൽ ഹെയർഷിപ്പ് /റിലേഷൻഷിപ്പ് സർട്ടിഫിക്കറ്റ്"),
    ("student_school_certificate", "വിദ്യാഭ്യാസ ആനുകൂല്യത്തിന് വിദ്യാർത്ഥിയുടെ സ്കൂൾ സർട്ടിഫിക്കറ്റ്"),
    ("funeral_bill", "ശവസംസ്കാര ചെലവ് ലഭ്യമാകുന്നതിന് ആയതിന്റെ ബിൽ"),
    ("disability_certificate", "ആശുപത്രിയിലെ കേസ് സർട്ടിഫിക്കറ്റ്/ഡിസെബിലിറ്റി സർട്ടിഫിക്കറ്റ്"),
    ("doctor_certificate_form", "ഡോക്ടറുടെ സർട്ടിഫിക്കറ്റ് നിശ്ചിത ഫോറത്തിൽ"),
    ("police_report", "പോലീസ് സ്റ്റേഷനിൽ നിന്നുളള റിപ്പോർട്ട്"),
    ("collector_rdo_certificate", "കളക്ടർ/ആർ.ഡി.ഒ നൽകുന്ന സർട്ടിഫിക്കററ്"),
    ("hospital_discharge_card", "ആശുപത്രി ഡിസ്ചാർജ്ജ് കാർഡ്"),
    ("hospital_bill", "ആശുപത്രിയിലെ അസ്സൽ ബിൽ"),
])
apply_for("FISH-01", [dict(scheme_id="FISH-01", type_id=FISH_OFFICE, mode="in_person",
    note_en="Report the accident to the police station and the Matsyaboard Fisheries Officer at once. Submit the claim in 4 copies to the Fisheries Officer.",
    note_ml="അപകടമുണ്ടായാലുടനെ പോലീസ് സ്റ്റേഷനിലും മത്സ്യബോർഡ് ഫിഷറീസ് ഓഫീസറെയും അറിയിക്കണം. അപേക്ഷ (നാല് പകർപ്പുകൾ) ഫിഷറീസ് ഓഫീസർക്ക് നൽകണം.",
    source_id=K, quote="കഴിയുന്നത്ര വേഗം അപേക്ഷ തയ്യാറാക്കി (നാല് പകർപ്പുകൾ) ഫിഷറീസ് ഓഫീസർക്ക് നൽകണം.")])

# ---------------------------------------------------------------- FISH-06 old-age pension (2025) — upgraded from partial
scheme("FISH-06", name_en="Old Age Pension for Fishermen", name_ml="വാർദ്ധക്യകാല പെൻഷൻ",
       authority_en="Kerala Fishermen's Welfare Fund Board", authority_ml="കേരള മത്സ്യത്തൊഴിലാളി ക്ഷേമനിധി ബോർഡ്",
       summary_en="Monthly pension of ₹1,600 for Board members who have retired from fishing after age 60.",
       summary_ml="60 വയസ്സ് പൂർത്തിയാക്കി തൊഴിലിൽ നിന്ന് വിരമിച്ച ബോർഡ് അംഗങ്ങൾക്ക് പ്രതിമാസം 1,600 രൂപ പെൻഷൻ.",
       eligibility_verification="verified", documents_verification="verified", apply_verification="verified",
       notes="Upgraded from partial using the 2025 KFWFB guideline p.18-19. CONFLICT: the older department page said Rs 450/month for 'registered fishermen/widows'; the 2025 guideline gives Rs 1,600 and the conditions below. Widows are now covered by FISH-10.")
delete_where(C, scheme_id="FISH-06", condition_id="c1")
Q5 = "കേരളത്തിൽ കുറഞ്ഞത് 10 വർഷമെങ്കിലും മത്സ്യത്തൊഴിലാളിയായി ജോലി ചെയ്ത് ഉപജീവനം കഴിക്കുകയും 60 വയസ്സ് പൂർത്തിയാക്കുകയും തൊഴിലിൽ നിന്നും വിരമിക്കുകയും ചെയ്യുന്ന ക്ഷേമനിധി അംഗങ്ങൾ പെൻഷൻ ലഭിക്കുന്നതിന് അർഹരാണ്."
cond("FISH-06", "g0", 0, "livelihood", "includes", "fishing", "വാർദ്ധക്യകാലത്ത് വരുമാന മാർഗ്ഗമില്ലാതെയാവുന്ന മത്സ്യത്തൊഴിലാളികൾക്ക് പ്രതിമാസ പെൻഷൻ നൽകുന്നതിനുള്ള പദ്ധതിയാണിത്.", "p.18, scheme 5")
cond("FISH-06", "c1", 1, "kfwfb_member", "eq", "yes", Q5, "p.18, eligibility", changeable="yes", how_en=HOW_KFWFB_EN, how_ml=HOW_KFWFB_ML)
cond("FISH-06", "c2", 2, "age", "gte", 60, Q5, "p.18, eligibility")
cond("FISH-06", "c3", 3, "fisher_work_10_years", "eq", "yes", Q5, "p.18, eligibility")
cond("FISH-06", "c4", 4, "retired_from_fishing", "eq", "yes", Q5, "p.18, eligibility")
cond("FISH-06", "c5", 5, "kfwfb_membership_years", "gte", 5, "പെൻഷൻ അപേക്ഷ തിയ്യതി തൊട്ട് 5 വർഷം മുമ്പെങ്കിലും മത്സ്യബോർഡിൽ അംഗത്വം എടുത്ത് വിഹിതം അടച്ചിരിക്കണം.", "p.18, eligibility")
docs_for("FISH-06", [
    ("age_proof", "അപേക്ഷന്റെ വയസ്സ് തെളിയിക്കുന്നതിന് സ്കൂളിൽ നിന്നോ, ജനന മരണ രജിസ്ട്രാറിൽ നിന്നോ, ലഭിക്കുന്ന സർട്ടിഫിക്കററ്, പാസ്സ്പോർട്ട്, ഡ്രൈവിംഗ് ലൈസൻസ് ഇവയൊന്നുമില്ലെങ്കിൽ നിശ്ചിത ഫോറത്തിൽ ഉള്ള മെഡിക്കൽ സർട്ടിഫിക്കറ്റ്."),
    ("kfwfb_passbook", "ഫിഷറീസ് ഓഫീസർ സാക്ഷ്യപ്പെടുത്തിയ അപേക്ഷകന്റെ മത്സ്യബോർഡ് പാസ്സ് ബുക്കിന്റെ പകർപ്പ്"),
    ("bank_passbook_aadhaar_linked", "ആധാർ ലിങ്ക് ചെയ്ത വ്യക്തിഗത ബാങ്ക് അക്കൗണ്ട് പാസ് ബുക്കിന്റെ പകർപ്പ്"),
    ("ration_card", "റേഷൻ കാർഡിന്റെ പകർപ്പ് ഫിഷറീസ് ഓഫീസർ സാക്ഷ്യപ്പെടുത്തിയത്."),
    ("aadhaar_copy", "ആധാർ കാർഡിന്റെ പകർപ്പ്"),
    ("income_certificate", "വരുമാന സർട്ടിഫിക്കറ്റ്"),
])
apply_for("FISH-06", [dict(scheme_id="FISH-06", type_id=FISH_OFFICE, mode="online",
    note_en="Apply online on www.fims.kerala.gov.in to the Fisheries Officer once the person completes 60.",
    note_ml="60 വയസ്സ് പൂർത്തിയാക്കുന്ന മുറക്ക് www.fims.kerala.gov.in പോർട്ടൽ മുഖേന ഫിഷറീസ് ഓഫീസർക്ക് അപേക്ഷ സമർപ്പിക്കുക.",
    source_id=K, quote="60 വയസ്സ് പൂർത്തിയാക്കുന്ന മുറക്ക് അപേക്ഷ www.fims.kerala.gov.in എന്ന പോർട്ടൽ മുഖേന ഫിഷറീസ് ഓഫീസർക്ക് സമർപ്പിക്കണം.")])

# ---------------------------------------------------------------- FISH-10 widow pension
scheme("FISH-10", name_en="Widow Pension (Fishermen's Welfare Fund)", name_ml="വിധവ പെൻഷൻ",
       authority_en="Kerala Fishermen's Welfare Fund Board", authority_ml="കേരള മത്സ്യത്തൊഴിലാളി ക്ഷേമനിധി ബോർഡ്", category="fishing",
       summary_en="Monthly pension of ₹1,600 for the wife of a Board member who died before receiving any old-age pension instalment.",
       summary_ml="വാർദ്ധക്യകാല പെൻഷൻ ഒരു ഗഡുപോലും കൈപ്പറ്റുന്നതിനുമുമ്പ് മരണമടഞ്ഞ ബോർഡ് അംഗങ്ങളുടെ ഭാര്യമാർക്ക് പ്രതിമാസം 1,600 രൂപ.",
       eligibility_verification="verified", documents_verification="verified", apply_verification="verified",
       notes="2025 KFWFB guideline p.30-31. The 90-day application limit is encoded as a condition because the page states it. Condition 3 on p.30 is printed blank.")
cond("FISH-10", "g0", 0, "livelihood", "includes", "fishing", "മത്സ്യത്തൊഴിലാളികളുടെ മരണം മൂലം നിരാലംബരായി തീരുന്ന അവരുടെ ഭാര്യമാർക്ക് ആശ്വാസ ധനസഹായമെന്ന നിലയിൽ പ്രതിമാസ പെൻഷൻ നൽകുന്നതിനുള്ള പദ്ധതിയാണിത്.", "p.30")
cond("FISH-10", "c1", 1, "kfwfb_family_status", "includes", "widow_of_fisher", "മത്സ്യബോർഡിൽ വിഹിതമടച്ച് അംഗത്വമെടുത്ത ശേഷം വാർദ്ധക്യകാല പെൻഷൻ പദ്ധതി പ്രകാരമുള്ള പെൻഷൻ തുക ഒരു ഗഡുപോലും കൈപ്പറ്റുന്നതിനുമുമ്പായി മരണമടയുന്ന മത്സ്യത്തൊഴിലാളികളുടെ ഭാര്യമാർക്ക് ഈ പദ്ധതി പ്രകാരം പ്രതിമാസ പെൻഷൻ തുക 1600/- രൂപ നൽകി വരുന്നു.", "p.30")
cond("FISH-10", "c2", 2, "regular_income_job", "eq", "no", "സ്ഥിര വരുമാനമുള്ള ഉദ്യോഗമുണ്ടെങ്കിൽ ഈ പദ്ധതി പ്രകാരമുള്ള പെൻഷൻ ലഭിക്കുകയില്ല.", "p.30, condition 1")
cond("FISH-10", "c3", 3, "widow_remarried", "eq", "no", "പുനർവിവാഹം നടക്കുന്ന പക്ഷം വിവാഹം നടക്കുന്ന മാസം മുതൽ പെൻഷൻ റദ്ദുചെയ്യുന്നതാണ്.", "p.30, condition 2")
cond("FISH-10", "c4", 4, "death_within_3_months", "eq", "yes", "മരണം നടന്ന് 90 ദിവസത്തിനുള്ളിൽ ഫിഷറീസ് ഓഫീസർക്ക് അപേക്ഷ നൽകണം.", "p.30, how to apply")
docs_for("FISH-10", [
    ("member_death_certificate_attested", "മത്സ്യത്തൊഴിലാളിയുടെ മരണ സർട്ടിഫിക്കറ്റ് പകർപ്പ് ഫിഷറീസ് ഓഫീസർ സാക്ഷ്യപ്പെടുത്തിയത്."),
    ("passbook_contribution_pages", "മത്സ്യബോർഡ് വിഹിതമടച്ച പേജിന്റെയും പാസ്സ് ബുക്കിന്റെയും പൂർണ്ണമായ പകർപ്പ് ഫിഷറീസ് ഓഫീസർ സാക്ഷ്യപ്പെടുത്തിയത്."),
    ("no_other_pension_proof", "മറ്റു പെൻഷൻ ലഭിക്കുന്നില്ല എന്ന് തെളിയിക്കുന്ന രേഖ"),
    ("aadhaar_copy", "ആധാർ കാർഡിന്റെ പകർപ്പ്"),
    ("bank_passbook_aadhaar_linked", "വ്യക്തിഗത ബാങ്ക് പാസ് ബുക്കിന്റെ പകർപ്പ് (ആധാർ ലിങ്ക് ചെയ്ത ബാങ്ക് അക്കൗണ്ട്)"),
    ("ration_card_family_page", "റേഷൻ കാർഡിന്റെ പകർപ്പ് (കുടുംബ വിവരം അടങ്ങിയ പേജ് സഹിതം)"),
    ("not_remarried_certificate", "പുനർവിവഹം നടത്തിയിട്ടില്ലെന്നും സ്ഥിരവരുമാനമുള്ള തൊഴിൽ ഇല്ലെന്നുംതെളിയിക്കുന്നതിന് പഞ്ചായത്ത്, മുൻസിപ്പാലിറ്റി/ കോർപ്പറേഷൻ/ സെക്രട്ടറി വാർഡ് കൗൺസിലർ നൽകുന്ന സാക്ഷ്യപത്രം"),
])
apply_for("FISH-10", [dict(scheme_id="FISH-10", type_id=FISH_OFFICE, mode="in_person",
    note_en="Submit 2 copies of the prescribed form to the Fisheries Officer in charge of the village where the fisherman was a member, within 90 days of the death.",
    note_ml="നിശ്ചിത ഫോറത്തിൽ 2 കോപ്പികൾ, മരിച്ച മത്സ്യത്തൊഴിലാളി അംഗമായിരുന്ന ഗ്രാമത്തിന്റെ ചുമതല വഹിക്കുന്ന ഫിഷറീസ് ഓഫീസർക്ക്, മരണം നടന്ന് 90 ദിവസത്തിനുള്ളിൽ സമർപ്പിക്കുക.",
    source_id=K, quote="നിശ്ചിത ഫോറത്തിൽ അപേക്ഷയുടെ 2 കോപ്പികൾ തയ്യാറാക്കി മരിച്ച മത്സ്യത്തൊഴിലാളി അംഗമായിരുന്ന ഗ്രാമത്തിന്റെ ചുമതല വഹിക്കുന്ന ഫിഷറീസ് ഓഫീസർക്ക് സമർപ്പിക്കേണ്ടതാണ്.")])

# ---------------------------------------------------------------- FISH-11 marriage assistance
scheme("FISH-11", name_en="Marriage Assistance for Fishermen's Daughters", name_ml="വിവാഹ ധനസഹായ പദ്ധതി",
       authority_en="Kerala Fishermen's Welfare Fund Board", authority_ml="കേരള മത്സ്യത്തൊഴിലാളി ക്ഷേമനിധി ബോർഡ്", category="fishing",
       summary_en="₹25,000 towards the marriage of a fisherman's daughter.", summary_ml="മത്സ്യത്തൊഴിലാളികളുടെ പെൺമക്കളുടെ വിവാഹത്തിന് 25,000 രൂപ ധനസഹായം.",
       eligibility_verification="verified", documents_verification="verified", apply_verification="verified",
       notes="2025 KFWFB guideline p.15-16. Membership duration is 3 years from the first contribution and the application window is 60 days (the English summary said 60 months / 3 months; the page says otherwise). Condition 5 on p.15 (members removed from the list after an insured disability remain eligible) is not modelled. For the widow category the contributions must be paid up to the husband's death.")
cond("FISH-11", "g0", 0, "livelihood", "includes", "fishing", "മത്സ്യത്തൊഴിലാളികളുടെ പെൺമക്കളുടെ വിവാഹത്തിന് ധനസഹായം നൽകുന്നതിനുള്ള പദ്ധതിയാണിത്.", "p.15")
QM = "60 വയസ്സിൽ താഴെ പ്രായമുള്ള മത്സ്യത്തൊഴിലാളി; മത്സ്യത്തൊഴിലാളികളുടെ വിധവകൾ; 60 വയസ്സിൽ കൂടുതൽ പ്രായമുള്ള മത്സ്യത്തൊഴിലാളികൾ; പെൻഷൻകാർ; പെൻഷൻകാരുടെ വിധവകൾ; വധു മത്സ്യത്തൊഴിലാളിയാണെങ്കിൽ; മത്സ്യത്തൊഴിലാളിയായ പിതാവ്/മാതാവ് മരണപ്പെട്ടാൽ"
cond("FISH-11", "c1", 1, "kfwfb_family_status", "includes", "member_under_60|widow_of_fisher|member_over_60|pensioner|widow_of_pensioner|bride_is_fisher|fisher_parent_died", QM, "p.15, eligibility items 1-7 (joined)")
cond("FISH-11", "c2a", 2, "kfwfb_family_status", "includes", "member_under_60|widow_of_fisher|bride_is_fisher|fisher_parent_died", QM, "p.15, items 1-2 and 6-7 have no one-time limit")
cond("FISH-11", "c2b", 2, "prior_marriage_assistance", "eq", "no", "മുൻപൊരിക്കലും ഈ സഹായം ലഭിച്ചിട്ടില്ലെങ്കിൽ ഒരു മകളുടെ വിവാഹത്തിന്", "p.15, items 3-5")
cond("FISH-11", "c3", 3, "daughter_marriage_60_days", "eq", "yes", "വിവാഹ തിയ്യതി മുതൽ 60 ദിവസത്തിനകം അപേക്ഷ ഫിഷറീസ് ഓഫീസർക്ക് സമർപ്പിക്കേണ്ടതാണ്.", "p.16, how to apply 1")
cond("FISH-11", "c4", 4, "bride_18", "eq", "yes", "വധുവിന് വിവാഹ തിയ്യതിക്ക് 18 വയസ്സ് പൂർത്തിയായിരിക്കണം", "p.15, condition 1")
cond("FISH-11", "c5", 5, "annual_family_income", "lt", 50000, "അപേക്ഷകരുടെ കുടുംബ വാർഷിക വരുമാനം 50,000/- രൂപയിൽ താഴെ ആയിരിക്കണം", "p.15, condition 2")
cond("FISH-11", "c6", 6, "kfwfb_contributions_paid", "eq", "yes", "മത്സ്യത്തൊഴിലാളി വിഹിതങ്ങൾ പൂർണ്ണമായും അടച്ചിരിക്കണം", "p.15, condition 3", changeable="yes", how_en=HOW_CONTRIB_EN, how_ml=HOW_CONTRIB_ML)
cond("FISH-11", "c7", 7, "kfwfb_membership_years", "gte", 3, "അംഗത്വമെടുത്ത് ആദ്യമായി വിഹിതം അടച്ചതുമുതൽ 3 വർഷം കഴിഞ്ഞാൽ മാത്രമേ ഈ പദ്ധതി പ്രകാരമുള്ള ധനസഹായം ലഭിക്കുകയുള്ളൂ.", "p.15, condition 4")
docs_for("FISH-11", [
    ("bride_age_proof", "വധുവിന്റെ വയസ്സ് തെളിയിക്കുന്ന രേഖ"),
    ("income_proof_ration_or_certificate", "50,000/- രൂപയിൽ താഴെ റേഷൻകാർഡിൽ വാർഷിക വരുമാനം രേഖപ്പെടുത്തിയതിന്റെ പകർപ്പ് ( റേഷൻ കാർഡ്/വരുമാന സർട്ടിഫിക്കറ്റ്)"),
    ("kfwfb_passbook", "മത്സ്യബോർഡ് പാസ് ബുക്കിന്റെ പകർപ്പ് ( ഫിഷറീസ് ഓഫീസർ സാക്ഷ്യപ്പെടുത്തിയത്)"),
    ("marriage_certificate", "തദ്ദേശസ്വയംഭരണ സ്ഥാപനങ്ങളിൽ നിന്നും ലഭിക്കുന്ന വിവാഹ സർട്ടിഫിക്കറ്റ്"),
    ("aadhaar_copy", "അപേക്ഷകന്റെ/അപേക്ഷകയുടെ ആധാർ നമ്പർ"),
])
apply_for("FISH-11", [dict(scheme_id="FISH-11", type_id=FISH_OFFICE, mode="in_person",
    note_en="Submit 2 copies of the prescribed form to the Fisheries Officer within 60 days of the marriage.",
    note_ml="വിവാഹ തിയ്യതി മുതൽ 60 ദിവസത്തിനകം നിശ്ചിത ഫോറത്തിൽ 2 പകർപ്പുകൾ ഫിഷറീസ് ഓഫീസർക്ക് സമർപ്പിക്കുക.",
    source_id=K, quote="വിവാഹ തിയ്യതി മുതൽ 60 ദിവസത്തിനകം അപേക്ഷ ഫിഷറീസ് ഓഫീസർക്ക് സമർപ്പിക്കേണ്ടതാണ്.")])

# ---------------------------------------------------------------- death-related schemes
Q_U60 = "മരിച്ച മത്സ്യത്തൊഴിലാളി 60 വയസ്സിൽ താഴെ പ്രായമുള്ള ആളായിരിക്കണം."
Q_A60 = "60 വയസ്സിനു ശേഷം തൊഴിൽ തുടരുകയും പട്ടികയിൽ പേരുണ്ടായിരിക്കുകയും വിഹിതങ്ങൾ കൃത്യമായി അടക്കുകയും ചെയ്യുന്നവരുടെ ആശ്രിതർക്ക് ഈ പദ്ധതിയുടെ ആനുകൂല്യം ലഭിക്കുന്നതാണ്."
Q_PAID = "ക്ഷേമനിധി വിഹിതം പൂർണ്ണമായി അടച്ചിരിക്കണം."

scheme("FISH-12", name_en="Assistance to Dependants on a Fisherman's Non-Accidental Death", name_ml="അപകടം കൊണ്ടല്ലാതെ ആകസ്മിക കാരണങ്ങളാൽ ഉണ്ടാകുന്ന മരണത്തിന് ആശ്രിതർക്ക് ധനസഹായം",
       authority_en="Kerala Fishermen's Welfare Fund Board", authority_ml="കേരള മത്സ്യത്തൊഴിലാളി ക്ഷേമനിധി ബോർഡ്", category="fishing",
       summary_en="₹1,00,000 to the dependants when a fisherman dies suddenly from a cause other than an accident, during or right after fishing, and the group insurance does not apply.",
       summary_ml="മത്സ്യബന്ധനസമയത്തോ തൊട്ടുപിന്നാലെയോ അപകടം കൊണ്ടല്ലാതെ ആകസ്മിക കാരണങ്ങളാൽ മരിക്കുന്ന മത്സ്യത്തൊഴിലാളികളുടെ ആശ്രിതർക്ക് 1,00,000 രൂപ.",
       eligibility_verification="verified", documents_verification="verified", apply_verification="verified",
       notes="2025 KFWFB guideline p.13-14. The 'during or right after fishing' timing is shown in the summary but not asked as a separate question.")
cond("FISH-12", "g0", 0, "livelihood", "includes", "fishing", "മത്സ്യബന്ധനസമയത്തോ, തൊട്ടുപിന്നാലെയോ അപകടം കൊണ്ടല്ലാതെ ആകസ്മിക കാരണങ്ങളാൽ മരിക്കുകയും , ഗ്രൂപ്പ് ഇൻഷ്വറൻസ് പദ്ധതി പ്രകാരമുള്ള ധനസഹായം ലഭിക്കാൻ അർഹതയില്ലാതാകുകയും ചെയ്യുന്ന മത്സ്യത്തൊഴിലാളികളുടെ ആശ്രിതർക്ക്", "p.13")
cond("FISH-12", "c1", 1, "member_death", "eq", "yes", "മത്സ്യബന്ധനസമയത്തോ, തൊട്ടുപിന്നാലെയോ അപകടം കൊണ്ടല്ലാതെ ആകസ്മിക കാരണങ്ങളാൽ മരിക്കുകയും , ഗ്രൂപ്പ് ഇൻഷ്വറൻസ് പദ്ധതി പ്രകാരമുള്ള ധനസഹായം ലഭിക്കാൻ അർഹതയില്ലാതാകുകയും ചെയ്യുന്ന മത്സ്യത്തൊഴിലാളികളുടെ ആശ്രിതർക്ക്", "p.13")
cond("FISH-12", "c2", 2, "death_by_accident", "eq", "no", "അപകടം കൊണ്ടല്ലാതെ ആകസ്മിക കാരണങ്ങളാൽ", "p.13")
cond("FISH-12", "c3a", 3, "deceased_under_60", "eq", "yes", Q_U60, "p.13, condition 1")
cond("FISH-12", "c3b", 3, "deceased_active_after_60", "eq", "yes", Q_A60, "p.13, condition 3")
cond("FISH-12", "c4", 4, "kfwfb_contributions_paid", "eq", "yes", "ക്ഷേമനിധി വിഹിതം പൂർണ്ണമായി അടച്ചിരിക്കണം.", "p.13, condition 2")
docs_for("FISH-12", [
    ("death_certificate", "മരണ സർട്ടിഫിക്കറ്റ് – അസ്സൽ"),
    ("doctor_certificate_cause", "ആശുപത്രിയിൽ പ്രവേശിച്ച ശേഷമാണ് മരണമെങ്കിൽമരണം സാധാരണ രോഗം കൊണ്ടുണ്ടായതല്ലെന്നും അപ്രതീക്ഷിതമായ കാരണങ്ങൾകൊണ്ടായിരുന്നുവെന്നും തെളിയിക്കുന്നതിന് ചികിൽസിച്ച ഡോക്ടറുടെ സർട്ടിഫിക്കറ്റ്"),
    ("police_report", "പോലീസ് സ്റ്റേഷനിൽ നിന്നും ലഭിക്കുന്ന റിപ്പോർട്ടിന്റെ പകർപ്പ്."),
    ("kfwfb_passbook", "മരിച്ചയാളിന്റെ മത്സ്യബോർഡ് പാസ്സ് ബുക്കിന്റെ ഫോട്ടോ കോപ്പി ഫിഷറീസ് ഓഫീസർ സാക്ഷ്യപ്പെടുത്തിയത്."),
    ("bank_details", "അവകാശിയുടെ ബാങ്ക് അക്കൌണ്ട് വിശദാംശങ്ങൾ"),
    ("aadhaar_copy", "അവകാശിയുടെ ആധാർ നമ്പർ"),
])
apply_for("FISH-12", [dict(scheme_id="FISH-12", type_id=FISH_OFFICE, mode="in_person",
    note_en="Inform the Fisheries Officer and the nearest police station at once. Submit 3 copies of the application with the documents to the Fisheries Officer.",
    note_ml="ഉടനെ തന്നെ ഫിഷറീസ് ഓഫീസറേയും തൊട്ടടുത്ത പോലീസ് സ്റ്റേഷനിലും അറിയിക്കണം. അപേക്ഷയുടെ മൂന്ന് പകർപ്പുകൾ രേഖകൾ സഹിതം ഫിഷറീസ് ഓഫീസർക്ക് സമർപ്പിക്കുക.",
    source_id=K, quote="സാമ്പത്തിക സഹായത്തിനുള്ള അപേക്ഷയുടെ മൂന്ന് പകർപ്പുകൾ തയ്യാറാക്കി താഴെ പറയുന്ന രേഖകൾ സഹിതം ബന്ധപ്പെട്ട ഫിഷറീസ് ഓഫീസർക്ക് സമർപ്പിക്കേണ്ടതാണ്.")])

scheme("FISH-13", name_en="Death Assistance to Fishermen's Dependants", name_ml="മത്സ്യത്തൊഴിലാളികളുടെ മരണത്തോടനുബന്ധിച്ച് ആശ്രിതർക്കുള്ള ധനസഹായ പദ്ധതി",
       authority_en="Kerala Fishermen's Welfare Fund Board", authority_ml="കേരള മത്സ്യത്തൊഴിലാളി ക്ഷേമനിധി ബോർഡ്", category="fishing",
       summary_en="₹15,000 to the dependants when a fisherman dies in any circumstance. Amounts already paid under another death-related scheme are deducted.",
       summary_ml="മത്സ്യത്തൊഴിലാളികൾ ഏതു സാഹചര്യത്തിൽ മരിച്ചാലും ആശ്രിതർക്ക് 15,000 രൂപ. മരണവുമായി ബന്ധപ്പെട്ട മറ്റു പദ്ധതി പ്രകാരം ലഭിച്ച തുക കുറയ്ക്കും.",
       eligibility_verification="verified", documents_verification="verified", apply_verification="verified",
       notes="2025 KFWFB guideline p.21-22. Condition 4 (Rs 1,000 to whoever paid funeral costs when there is no dependant) is not modelled.")
cond("FISH-13", "g0", 0, "livelihood", "includes", "fishing", "മത്സ്യത്തൊഴിലാളികൾ ഏതു സാഹചര്യത്തിൽ മരിച്ചാലും ആശ്രിതർക്ക് ധനസഹായം നൽകുന്ന പദ്ധതിയാണിത്.", "p.21")
cond("FISH-13", "c1", 1, "member_death", "eq", "yes", "മത്സ്യത്തൊഴിലാളികൾ ഏതു സാഹചര്യത്തിൽ മരിച്ചാലും ആശ്രിതർക്ക് ധനസഹായം നൽകുന്ന പദ്ധതിയാണിത്.", "p.21")
cond("FISH-13", "c2a", 2, "deceased_under_60", "eq", "yes", "മരിച്ച മത്സ്യത്തൊഴിലാളി 60 വയസ്സിൽ താഴെ പ്രായമുള്ള ആളായിരിക്കണം. ക്ഷേമനിധി വിഹിതം പൂർണ്ണമായി അടച്ചിരിക്കണം.", "p.21, condition 1")
cond("FISH-13", "c2b", 2, "deceased_active_after_60", "eq", "yes", "60 വയസ്സിനുശേഷം തൊഴിൽ തുടരുകയും പട്ടികയിൽ പേരുണ്ടായിരിക്കുകയും വിഹിതങ്ങൾ കൃത്യമായി അടക്കുകയും ചെയ്യുന്നവരുടെ ആശ്രിതർക്ക് ഈ പദ്ധതിയുടെ ആനുകൂല്യം ലഭിക്കുന്നതാണ്.", "p.21, condition 2")
cond("FISH-13", "c3", 3, "kfwfb_contributions_paid", "eq", "yes", "മരിച്ച മത്സ്യത്തൊഴിലാളി 60 വയസ്സിൽ താഴെ പ്രായമുള്ള ആളായിരിക്കണം. ക്ഷേമനിധി വിഹിതം പൂർണ്ണമായി അടച്ചിരിക്കണം.", "p.21, condition 1")
cond("FISH-13", "c4", 4, "death_within_3_months", "eq", "yes", "മത്സ്യത്തൊഴിലാളി മരണപ്പെട്ട് മൂന്ന് മാസത്തിനകം ഈ അപേക്ഷ സമർപ്പിക്കേണ്ടതാണ്.", "p.21, how to apply")
docs_for("FISH-13", [
    ("kfwfb_passbook", "മരിച്ചയാളിന്റെ മത്സ്യബോർഡ് പാസ്സ് ബുക്കിന്റെ സാക്ഷ്യപ്പെടുത്തിയ പകർപ്പും"),
    ("death_certificate", "മരണസർട്ടിഫിക്കറ്റും"),
    ("legal_heir_certificate", "ലീഗൽ ഹെയർഷിപ്പ് സർട്ടിഫിക്കറ്റും"),
    ("aadhaar_copy", "അപേക്ഷകന്റെ ആധാർ കാർഡിന്റെ പകർപ്പ്"),
    ("bank_passbook", "ബാങ്ക് പാസ്സ് ബുക്കിന്റെ പകർപ്പ്"),
])
apply_for("FISH-13", [dict(scheme_id="FISH-13", type_id=FISH_OFFICE, mode="in_person",
    note_en="Dependants submit 2 copies of the application to the Fisheries Officer within 3 months of the death, stating the relationship and the date of death.",
    note_ml="മരിച്ച് മൂന്ന് മാസത്തിനകം ആശ്രിതർ അപേക്ഷയുടെ 2 പകർപ്പുകൾ ബന്ധവും മരിച്ച തീയതിയും വ്യക്തമാക്കി ഫിഷറീസ് ഓഫീസർക്ക് സമർപ്പിക്കുക.",
    source_id=K, quote="ആശ്രിതർ അപേക്ഷയുടെ 2 പകർപ്പുകൾ തയ്യാറാക്കി ബന്ധപ്പെട്ട ഫിഷറീസ് ഓഫീസർക്ക് സമർപ്പിച്ചിരിക്കണം.")])

scheme("FISH-14", name_en="Funeral Assistance on the Death of a Fisherman's Dependant", name_ml="മത്സ്യത്തൊഴിലാളികളുടെ ആശ്രിതരുടെ മരണാനന്തര ചെലവുകൾക്ക് ധനസഹായം",
       authority_en="Kerala Fishermen's Welfare Fund Board", authority_ml="കേരള മത്സ്യത്തൊഴിലാളി ക്ഷേമനിധി ബോർഡ്", category="fishing",
       summary_en="₹2,000 towards funeral expenses when a fisherman's father, mother, spouse, minor son or unmarried daughter dies.",
       summary_ml="മത്സ്യത്തൊഴിലാളികളുടെ അച്ഛൻ, അമ്മ, ഭാര്യ/ഭർത്താവ്, മൈനർമാരായ ആൺമക്കൾ, അവിവാഹിതരായ പെൺമക്കൾ മരിക്കുമ്പോൾ ശവസംസ്കാര ചെലവിന് 2,000 രൂപ.",
       eligibility_verification="verified", documents_verification="verified", apply_verification="verified",
       notes="2025 KFWFB guideline p.17. Board membership is not stated as a condition on this page (the Matsyaboard passbook is a required document), so it is not encoded.")
Q4 = "മത്സ്യത്തൊഴിലാളികളുടെ അച്ഛൻ, അമ്മ, ഭാര്യ/ഭർത്താവ്, മൈനർമാരായ ആൺ മക്കൾ, അവിവാഹിതരായ പെൺമക്കൾ എന്നിവരുടെ മരണം സംഭവിക്കുമ്പോഴാണ് ഈ പദ്ധതി പ്രകാരമുള്ള ധനസഹായം നൽകുന്നത്."
cond("FISH-14", "g0", 0, "livelihood", "includes", "fishing", "മത്സ്യത്തൊഴിലാളികളുടെ ആശ്രിതർ മരണമടയുമ്പോൾ ശവസംസ്കാരം തുടങ്ങിയ ചെലവുകൾ നടത്തുന്നതിന് സാമ്പത്തിക സഹായം നൽകുന്ന പദ്ധതിയാണിത്.", "p.17")
cond("FISH-14", "c1", 1, "dependant_died", "eq", "yes", Q4 + " ആശ്രിതൻ മരിച്ച് മൂന്ന് മാസത്തിനകം നിശ്ചിത ഫോറത്തിൽ അപേക്ഷയുടെ 2 പകർപ്പുകൾ തയ്യാറാക്കി ഫിഷറീസ് ഓഫീസർക്ക് സമർപ്പിക്കണം.", "p.17")
docs_for("FISH-14", [
    ("dependant_death_certificate", "ആശ്രിതന്റെ മരണ സർട്ടിഫിക്കറ്റ്"),
    ("kfwfb_passbook", "മത്സ്യബോർഡ് പാസ്സ് ബുക്കിന്റെ പകർപ്പ് – ഫിഷറീസ് ഓഫീസർ സാക്ഷ്യപ്പെടുത്തിയത്."),
    ("legal_heir_certificate", "ലീഗൽ ഹെയർഷിപ്പ് / റിലേഷൻഷിപ്പ് സർട്ടിഫിക്കറ്റ്"),
])
apply_for("FISH-14", [dict(scheme_id="FISH-14", type_id=FISH_OFFICE, mode="in_person",
    note_en="Submit 2 copies of the prescribed form to the Fisheries Officer within 3 months of the death.",
    note_ml="മരിച്ച് മൂന്ന് മാസത്തിനകം നിശ്ചിത ഫോറത്തിൽ 2 പകർപ്പുകൾ ഫിഷറീസ് ഓഫീസർക്ക് സമർപ്പിക്കുക.",
    source_id=K, quote="ആശ്രിതൻ മരിച്ച് മൂന്ന് മാസത്തിനകം നിശ്ചിത ഫോറത്തിൽ അപേക്ഷയുടെ 2 പകർപ്പുകൾ തയ്യാറാക്കി ഫിഷറീസ് ഓഫീസർക്ക് സമർപ്പിക്കണം.")])

scheme("FISH-15", name_en="Relief for Temporary Disability due to Accident", name_ml="അപകടം മൂലമുണ്ടാകുന്ന താൽക്കാലിക അവശതയ്ക്ക് ആശ്വാസ ധനസഹായ പദ്ധതി",
       authority_en="Kerala Fishermen's Welfare Fund Board", authority_ml="കേരള മത്സ്യത്തൊഴിലാളി ക്ഷേമനിധി ബോർഡ്", category="fishing",
       summary_en="Relief while a fisherman cannot work after an accident: ₹100 for the first 7 days, then ₹15 a day, up to ₹500.",
       summary_ml="അപകടം മൂലം ജോലി ചെയ്യാൻ കഴിയാത്ത കാലത്തേക്ക് ആദ്യത്തെ ഏഴ് ദിവസത്തേക്ക് 100 രൂപ, പിന്നീട് ഓരോ ദിവസത്തേക്കും 15 രൂപ, പരമാവധി 500 രൂപ.",
       eligibility_verification="verified", documents_verification="verified", apply_verification="verified",
       notes="2025 KFWFB guideline p.20. Board membership is not stated as a condition on this page (the applicant's Matsyaboard passbook is a required document), so it is not encoded.")
cond("FISH-15", "g0", 0, "livelihood", "includes", "fishing", "അപകടം മൂലം താൽക്കാലികമായി തൊഴിൽ ചെയ്യാൻ കഴിയാതെ വരുന്ന കാലത്തേക്ക് ഉപജീവനത്തിനുള്ള ആശ്വാസ ധനസഹായം നൽകുന്ന പദ്ധതിയാണിത്.", "p.20")
cond("FISH-15", "c1", 1, "accident_unable_7_days", "eq", "yes", "കുറഞ്ഞത് ഏഴ് ദിവസമെങ്കിലും തൊഴിൽ ചെയ്യാൻ പറ്റാതെ വന്നാൽ മാത്രമെ ഈ പദ്ധതിയുടെ ആനുകൂല്യം ലഭിക്കുകയുള്ളൂ.", "p.20")
cond("FISH-15", "c2", 2, "accident_within_3_months", "eq", "yes", "അപകടം സംഭവിച്ച തിയ്യതി മുതൽ 3 മാസത്തിനകം ബന്ധപ്പെട്ട ഫിഷറീസ് ഓഫീസർക്ക് സമർപ്പിച്ചിരിക്കണം.", "p.20, how to apply")
docs_for("FISH-15", [
    ("govt_medical_certificate", "ചികിൽസിച്ച സർക്കാർ ഡോക്ടറിൽ നിന്നും ലഭിച്ച നിശ്ചിത ഫോറത്തിലുള്ള മെഡിക്കൽ സർട്ടിഫിക്കറ്റ്."),
    ("kfwfb_passbook", "അപേക്ഷകന്റെ മത്സ്യബോർഡ് പാസ് ബുക്കിന്റെ ശരി പകർപ്പ് (ഫിഷറീസ് ഓഫീസർ സാക്ഷ്യപ്പെടുത്തിയത്)"),
])
apply_for("FISH-15", [dict(scheme_id="FISH-15", type_id=FISH_OFFICE, mode="in_person",
    note_en="Inform the Fisheries Officer at once. Submit 2 copies of the prescribed form within 3 months of the accident.",
    note_ml="ഉടനെ ഫിഷറീസ് ഓഫീസറെ അറിയിക്കണം. അപകടം സംഭവിച്ച തിയ്യതി മുതൽ 3 മാസത്തിനകം നിശ്ചിത ഫോറത്തിൽ 2 പകർപ്പുകൾ സമർപ്പിക്കുക.",
    source_id=K, quote="ധനസഹായത്തിനുള്ള അപേക്ഷ നിശ്ചിത ഫോറത്തിൽ 2 പകർപ്പുകൾ തയ്യാറാക്കി അപകടം സംഭവിച്ച തിയ്യതി മുതൽ 3 മാസത്തിനകം ബന്ധപ്പെട്ട ഫിഷറീസ് ഓഫീസർക്ക് സമർപ്പിച്ചിരിക്കണം.")])

scheme("FISH-16", name_en="SSLC Cash Award for Fishermen's Children", name_ml="എസ്.എസ്.എൽ.സി പരീക്ഷയിൽ ഉന്നത വിജയം നേടുന്നവർക്ക് ക്യാഷ് അവാർഡ്",
       authority_en="Kerala Fishermen's Welfare Fund Board", authority_ml="കേരള മത്സ്യത്തൊഴിലാളി ക്ഷേമനിധി ബോർഡ്", category="fishing",
       summary_en="Cash award for fishermen's children with top SSLC results: ₹5,000 for A+ in all subjects, ₹4,000 for A+ in 9, ₹3,000 for A+ in 8.",
       summary_ml="എസ്.എസ്.എൽ.സി-യിൽ ഉന്നത വിജയം നേടുന്ന മത്സ്യത്തൊഴിലാളികളുടെ മക്കൾക്ക്: എല്ലാ വിഷയങ്ങളിലും എ+ 5,000 രൂപ, 9 വിഷയങ്ങളിൽ 4,000, 8 വിഷയങ്ങളിൽ 3,000.",
       eligibility_verification="verified", documents_verification="verified", apply_verification="partial",
       notes="2025 KFWFB guideline p.22-23. The separate category for the top 3 students of each Regional Fisheries Technical School (Rs 3,000) is not modelled. Applications open only after the Commissioner's notification each year.")
cond("FISH-16", "g0", 0, "livelihood", "includes", "fishing", "വിദ്യാഭ്യാസ പുരോഗതി നേടുന്നതിന് മത്സ്യത്തൊഴിലാളികളുടെ മക്കളായ വിദ്യാർത്ഥികൾക്ക് പ്രോത്സാഹനം നൽകുന്നതിനാണ് ഈ പദ്ധതി തയ്യാറാക്കിയിട്ടുള്ളത് .", "p.22")
cond("FISH-16", "c1", 1, "child_sslc_8_aplus", "eq", "yes", "എസ്.എസ്.എൽ.സി പരീക്ഷയ്ക്ക് 8വിഷയങ്ങളിലും എ+ ഗ്രേഡ്ലഭിക്കാത്തവിദ്യാർത്ഥികളെ അവാർഡിനായി പരിഗണിക്കുന്നതല്ല.", "p.23, condition 2")
Q8 = "വിദ്യാർത്ഥിയുടെ രക്ഷിതാവിന് ക്ഷേമനിധി അംഗത്വമുണ്ടെങ്കിൽ പോലും മത്സ്യബന്ധനം മുഖ്യ ഉപജീവനമായി സ്വീകരിച്ചിട്ടുള്ള പ്രവർത്തനോന്മുഖ മത്സ്യത്തൊഴിലാളിയായി തുടരുന്നുവെന്ന് ഉറപ്പുവരുത്തിയശേഷം മാത്രമേ അപേക്ഷ പരിഗണനീയമാണെന്ന ശുപാർശ ഫിഷറീസ് ഓഫീസർ മേലാധികാരിക്ക് സമർപ്പിക്കുകയുള്ളൂ."
cond("FISH-16", "c2", 2, "kfwfb_member", "eq", "yes", Q8, "p.23, condition 1", changeable="yes", how_en=HOW_KFWFB_EN, how_ml=HOW_KFWFB_ML)
cond("FISH-16", "c3", 3, "active_fisher", "eq", "yes", Q8, "p.23, condition 1")
docs_for("FISH-16", [
    ("sslc_certificate", "വിദ്യാർത്ഥിയുടെ എസ്.എസ്.എൽ.സി സട്ടിഫിക്കറ്റിന്റെ ശരിപ്പകർപ്പ് (ഫിഷറീസ് ഓഫീസർ സാക്ഷ്യപ്പെടുത്തിയത്) വയസ്സും, മാർക്കും തെളിയിക്കുന്നതിനായി."),
    ("parent_kfwfb_passbook", "രക്ഷിതാക്കളുടെ മത്സ്യബോർഡ് പാസ്സ് ബുക്കിന്റെ ശരി പകർപ്പ് (ഫിഷറീസ് ഓഫീസർ സാക്ഷ്യപ്പെടുത്തിയത്)"),
    ("photos_2", "വിദ്യാർത്ഥിയുടെ പാസ്പോർട്ട് സൈസ് ഫോട്ടോ 2 എണ്ണം"),
    ("aadhaar_copy", "ആധാർകാർഡിന്റെ പകർപ്പ്"),
    ("bank_passbook", "ബാങ്ക് പാസ് ബുക്കിന്റെ പകർപ്പ്"),
])
apply_for("FISH-16", [dict(scheme_id="FISH-16", type_id=FISH_OFFICE, mode="in_person",
    note_en="Watch for the Commissioner's notification after SSLC results. Apply on plain paper to the office named in it before the deadline; late applications are not accepted.",
    note_ml="എസ്.എസ്.എൽ.സി ഫലത്തിന് ശേഷം കമ്മീഷണർ പുറപ്പെടുവിക്കുന്ന വിജ്ഞാപനം ശ്രദ്ധിക്കുക. നിശ്ചിത തിയ്യതിക്കകം വെള്ളക്കടലാസിൽ അപേക്ഷ സമർപ്പിക്കുക.",
    source_id=K, quote="വിദ്യാർത്ഥികൾ നേരിട്ടോ രക്ഷിതാവ് മുഖേനയോ നിശ്ചിത തിയ്യതിക്കകം അപക്ഷകൻ ബന്ധപ്പെട്ട ഓഫീസർക്ക് സമർപ്പിക്കണം. സമയപരിധിക്കുശേഷം ലഭിക്കുന്ന അപേക്ഷകൾ സ്വീകരിക്കുന്നതല്ല.")])

scheme("FISH-17", name_en="Treatment Assistance for Serious Illness", name_ml="മാരകരോഗ ചികിൽസാ പദ്ധതി",
       authority_en="Kerala Fishermen's Welfare Fund Board", authority_ml="കേരള മത്സ്യത്തൊഴിലാളി ക്ഷേമനിധി ബോർഡ്", category="fishing",
       summary_en="Treatment help for fishermen with serious illness: up to ₹50,000 each for heart disease, kidney disease, cancer and brain tumour; up to ₹12,000 for paralysis; up to ₹5,000 for mental illness that can be cured with treatment. Treatment must be at a government or co-operative hospital, or on referral.",
       summary_ml="മാരകരോഗങ്ങൾക്ക് ചികിൽസാ സഹായം: ഹൃദ്രോഗം, വൃക്കരോഗം, ക്യാൻസർ, തലച്ചോറിലെ ട്യൂമർ എന്നിവയ്ക്ക് പരമാവധി 50,000 രൂപ വീതം; തളർവാതം 12,000; ചികിൽസിച്ച് ഭേദമാക്കാൻ പറ്റുന്ന മാനസികരോഗം 5,000. സർക്കാർ/സഹകരണ ആശുപത്രിയിലെ ചികിൽസയ്ക്ക് (അല്ലെങ്കിൽ റഫർ ചെയ്താൽ).",
       eligibility_verification="verified", documents_verification="verified", apply_verification="verified",
       notes="2025 KFWFB guideline p.25-26. 'Between 23 and 60' is encoded as 23 to 60 inclusive; confirm whether 23 and 60 themselves are included. The hospital restriction is shown in the summary, not asked.")
cond("FISH-17", "g0", 0, "livelihood", "includes", "fishing", "മാരകരോഗങ്ങൾ പിടിപെട്ട് കഷ്ടത അനുഭവിക്കുന്ന മത്സ്യത്തൊഴിലാളികൾക്ക് വിദഗ്ദ ചികിൽസയ്ക്കുവേണ്ട ധനസഹായം നൽകുന്നതിനുള്ള പദ്ധതിയാണിത്.", "p.25")
cond("FISH-17", "c1", 1, "serious_illness", "eq", "yes", "മാരകരോഗങ്ങൾ പിടിപെട്ട് കഷ്ടത അനുഭവിക്കുന്ന മത്സ്യത്തൊഴിലാളികൾക്ക് വിദഗ്ദ ചികിൽസയ്ക്കുവേണ്ട ധനസഹായം നൽകുന്നതിനുള്ള പദ്ധതിയാണിത്.", "p.25, list of diseases")
cond("FISH-17", "c2", 2, "age", "gte", 23, "അപേക്ഷകർ 23 വയസ്സിനും 60 വയസ്സിനും ഇടയിൽ പ്രായമുള്ളവരായിരിക്കണം.", "p.25, condition 1")
cond("FISH-17", "c3", 3, "age", "lte", 60, "അപേക്ഷകർ 23 വയസ്സിനും 60 വയസ്സിനും ഇടയിൽ പ്രായമുള്ളവരായിരിക്കണം.", "p.25, condition 1")
Q10 = "മത്സ്യബോർഡിൽ അംഗത്വമുള്ളവരായിരിക്കണം. ആദ്യമായി വിഹിതം അടച്ച് 5 വർഷമെങ്കിലും പൂർത്തിയാക്കുകയും കുടിശ്ശികയില്ലാതെ വിഹിതങ്ങളെല്ലാം അടക്കുകയും ചെയ്തിരിക്കണം."
cond("FISH-17", "c4", 4, "kfwfb_member", "eq", "yes", Q10, "p.25, condition 2", changeable="yes", how_en=HOW_KFWFB_EN, how_ml=HOW_KFWFB_ML)
cond("FISH-17", "c5", 5, "kfwfb_membership_years", "gte", 5, Q10, "p.25, condition 2")
cond("FISH-17", "c6", 6, "kfwfb_contributions_paid", "eq", "yes", Q10, "p.25, condition 2", changeable="yes", how_en=HOW_CONTRIB_EN, how_ml=HOW_CONTRIB_ML)
cond("FISH-17", "c7", 7, "annual_family_income", "lte", 50000, "വാർഷിക വരുമാനം 50,000/- രൂപയിൽ കൂടുതലായിരിക്കരുത്.", "p.25, condition 3")
docs_for("FISH-17", [
    ("income_proof_ration_or_certificate", "വരുമാനം തെളിയിക്കുന്നതിന് റേഷൻകാർഡിന്റെ /വരുമാന സർട്ടിഫിക്കറ്റിന്റെ പകർപ്പ് - ഫിഷറീസ് ഓഫീസർ സാക്ഷ്യപ്പെടുത്തയത്."),
    ("govt_hospital_medical_certificate", "സർക്കാർ ആശുപത്രിയിൽ ചികിസിച്ച ഡോക്ടർ നൽകുന്ന മെഡിക്കൽ സർട്ടിഫിക്കറ്റ്."),
    ("referral_document", "സർക്കാർ ആശുപത്രിയിൽ നിന്നും സ്വകാര്യ ആശുപത്രിയിലേക്ക് റഫർ ചെയ്തിട്ടുണ്ടെങ്കിൽ ആയതിനുള്ള രേഖ."),
    ("treatment_bills", "ചികിൽസാ ചിലവിലേക്കുള്ള ബില്ലുകൾ ചികിൽസിച്ച ഡോക്ടർ സാക്ഷ്യപ്പെടുത്തി സമർപ്പിക്കണം."),
    ("patient_age_proof", "രോഗിയുടെ വയസ്സ് തെളിയിക്കുന്നതിന് സ്കൂൾ സർട്ടിഫിക്കറ്റ് അഥവാ ജനന മരണ രജിസ്ട്രാറുടെ സർട്ടിഫിക്കറ്റ് ഇല്ലെങ്കിൽ ജ്ഞാനസ്നാന സർട്ടിഫിക്കറ്റ്."),
    ("kfwfb_passbook", "മത്സ്യബോർഡ് പാസ്സ് ബുക്കിന്റെ ശരി പകർപ്പ് ഫിഷറീസ് ഓഫീസർ സാക്ഷ്യപ്പെടുത്തിയത്."),
    ("prescription_copy", "ചികിൽസിക്കുന്ന ഡോക്ടറുടെ പ്രിസ്ക്രിപ്ഷന്റെ പകർപ്പ്"),
])
apply_for("FISH-17", [dict(scheme_id="FISH-17", type_id=FISH_OFFICE, mode="in_person",
    note_en="The patient or a dependant submits 3 copies of the prescribed form to the Fisheries Officer.",
    note_ml="രോഗിയോ രോഗിയുടെ ആശ്രിതരോ നിശ്ചിത ഫോറത്തിൽ 3 പകർപ്പുകൾ ഫിഷറീസ് ഓഫീസർക്ക് സമർപ്പിക്കണം.",
    source_id=K, quote="നിശ്ചിത ഫോറത്തിലുള്ള അപേക്ഷയുടെ 3 പകർപ്പുകൾ തയ്യാറാക്കി ഫിഷറീസ് ഓഫീസർക്ക് സമർപ്പിക്കണം.")])

scheme("FISH-18", name_en="Compensation for Fishing Equipment Lost in a Calamity", name_ml="പ്രകൃതിക്ഷോഭം മൂലം മത്സ്യബന്ധന ഉപകരണങ്ങൾ നഷ്ടപ്പെട്ടവർക്കുള്ള നഷ്ടപരിഹാരം",
       authority_en="Fisheries Department, Government of Kerala", authority_ml="കേരള സർക്കാർ മത്സ്യബന്ധന വകുപ്പ്", category="fishing",
       summary_en="Compensation when fishing equipment is lost or damaged in a sea calamity. We have the official application form, but not the rule that decides who qualifies.",
       summary_ml="കടൽക്ഷോഭം പോലുള്ള പ്രകൃതിക്ഷോഭത്തിൽ മത്സ്യബന്ധന ഉപകരണങ്ങൾ നഷ്ടപ്പെട്ടാൽ നഷ്ടപരിഹാരം. ഔദ്യോഗിക അപേക്ഷാ ഫോറം ലഭ്യമാണ്; അർഹത നിർണ്ണയിക്കുന്ന നിയമം ലഭ്യമല്ല.",
       eligibility_verification="informational", documents_verification="partial", apply_verification="partial",
       notes="Form only (3B, scanned). Display-only until the governing order is found.")
docs_for("FISH-18", [("calamity_application", None)], source="SRC-FISH-014")
apply_for("FISH-18", [dict(scheme_id="FISH-18", type_id=FISH_OFFICE, mode="in_person",
    note_en="Ask at the Fisheries Department district office; the form has an enquiry section for the Fisheries Department / Board.",
    note_ml="ഫിഷറീസ് വകുപ്പ് ജില്ലാ ഓഫീസിൽ അന്വേഷിക്കുക.", source_id="SRC-FISH-014", quote=None)])

# FISH-02: form field document ; FISH-07: form
for did, quote, srcid in [("scr_contribution_receipt", "അപേക്ഷകൻ അംഗമായ മത്സ്യത്തൊഴിലാളി സംഘത്തിന്റെ പേരും അംഗനമ്പരും നടപ്പു വർഷ ക്ഷേമനിധി വിഹിതം ഒടുക്കിയതിന്റെ രസീത് നമ്പരും", "SRC-FISH-012")]:
    delete_where(SD, scheme_id="FISH-02", doc_id=did)
    append(SD, dict(scheme_id="FISH-02", doc_id=did, source_id=srcid, quote=quote))
r = find(SC, scheme_id="FISH-02")[0]
setv(SC, r, "notes", (get(SC, r, "notes") or "") + " The collected application form (SRC-FISH-012) also declares no ownership of mechanised boats or beach-landing craft and no regular-income job in the family; these are NOT encoded because the form is an old blank-year form and no current guideline confirms them.")

# ---------------------------------------------------------------- plantation
upsert(LT, "type_id", dict(type_id="plantation_district_office", label_en="Small Plantation Workers' Welfare Fund — District Executive Office",
                           label_ml="ചെറുകിട തോട്ടം തൊഴിലാളി ക്ഷേമനിധി — ജില്ലാ എക്സിക്യൂട്ടീവ് ഓഫീസ്", scope="district"))
r = find(LT, type_id="spwwfb_office")[0]
setv(LT, r, "label_en", "Small Plantation Workers' Welfare Fund — Chief Executive Officer")
setv(LT, r, "label_ml", "ചെറുകിട തോട്ടം തൊഴിലാളി ക്ഷേമനിധി — ചീഫ് എക്സിക്യൂട്ടീവ് ഓഫീസർ")
upsert(L, "location_id", dict(location_id="LOC-PLNT-001", type_id="spwwfb_office", name_en="Office of the Chief Executive Officer, Kerala Small Plantation Workers' Welfare Fund",
       name_ml="ചീഫ് എക്സിക്യൂട്ടീവ് ഓഫീസറുടെ കാര്യാലയം, കേരള ചെറുകിട തോട്ടം തൊഴിലാളി ക്ഷേമനിധി", district="Kottayam", area="Azad Lane",
       address="Azad Lane, Kottayam - 686 001", phone="0481-2566672", source_id="SRC-PLNT-007", serves_districts=None,
       jurisdiction="Inspector of Plantation Nedumangad to Inspector of Plantation Mananthavady"))
pl = [
    ("LOC-PLNT-002", "Nedumangad", "Thiruvananthapuram", "Revenue Tower, Nedumangad PO, Thiruvananthapuram - 695 541", "0472-2802032", "Kollam",
     "Revenue District Thiruvananthapuram and taluks of Kollam, Karunagapally & Kottarakkara in Kollam District"),
    ("LOC-PLNT-003", "Pathanamthitta", "Pathanamthitta", "College Road, Excel Complex, Pathanamthitta PO - 689 645", "0468-2223069", "Alappuzha|Kottayam",
     "Revenue Districts Pathanamthitta, Alappuzha and Kottayam"),
    ("LOC-PLNT-004", "Pathanapuram", "Kollam", "Mini Civil Station, Pallimukku, Pathanapuram - 689 695", "0475-2352551", None,
     "Taluks of Pathanapuram & Kunnathur in Kollam District"),
    ("LOC-PLNT-005", "Peerumade", "Idukki", "Peerumade PO, Idukki - 685 531", "04869-233878", None, "Peerumade Taluk, Idukki District"),
    ("LOC-PLNT-006", "Vandanmedu", "Idukki", "Old Taluk Office Complex, Udumbanchola, Nedumkandam PO - 685 553", "04868-288386", None,
     "Vandanmedu Taluk, Idukki District (as printed)"),
    ("LOC-PLNT-007", "Munnar", "Idukki", "Devikulam PO, Idukki - 685 612", "04865-232565", None, "Devikulam Taluk, Idukki District"),
    ("LOC-PLNT-008", "Aluva", "Ernakulam", "Cart Stand Building, Market Road, opp. Fire Station, Aluva PO, Ernakulam - 683 101", "0484-2620284", "Thrissur|Idukki",
     "Thodupuzha Taluk (Idukki), Revenue District Ernakulam, Revenue District Thrissur"),
    ("LOC-PLNT-009", "Nemmara", "Palakkad", "Guruvayoorappan Building, Guru Complex, Ayiloor Road, Nemmara PO, Palakkad - 678 508", "04923-244070", None,
     "Revenue District Palakkad"),
    ("LOC-PLNT-010", "Manjeri", "Malappuram", "Karuvambram PO, Melakkam, Manjeri, Malappuram - 676 123", "04832-760204", "Kozhikode",
     "Revenue Districts Malappuram and Kozhikode (phone as printed: 04832-760204)"),
    ("LOC-PLNT-011", "Kalpetta", "Wayanad", "Asoothrana Bhavan, 4th Floor, Kalpetta - 673 122", "04936-204646", None,
     "Vythiri and Sulthan Bathery Taluks in Wayanad District"),
    ("LOC-PLNT-012", "Mananthavady", "Wayanad", "Mini Civil Station, Mananthavady PO, Wayanad - 670 645", "04935-241072", "Kannur|Kasaragod",
     "Mananthavady Taluk (Wayanad), Revenue Districts Kannur and Kasaragod"),
]
for lid, place, dist, addr, ph, serves, jur in pl:
    upsert(L, "location_id", dict(location_id=lid, type_id="plantation_district_office",
           name_en=f"Office of the District Executive Officer, {place}", name_ml=f"ജില്ലാ എക്സിക്യൂട്ടീവ് ഓഫീസറുടെ കാര്യാലയം, {place}",
           district=dist, area=place, address=addr, phone=ph, source_id="SRC-PLNT-007", serves_districts=serves, jurisdiction=jur))
for sid in ["PLNT-01", "PLNT-02", "PLNT-03", "PLNT-04"]:
    delete_where(A, scheme_id=sid, type_id="plantation_district_office")
    append(A, dict(scheme_id=sid, type_id="plantation_district_office", mode="in_person",
                   note_en="The District Executive Officer for your area handles the Fund's members. The application procedure for this benefit is not published in our sources.",
                   note_ml="നിങ്ങളുടെ പ്രദേശത്തെ ജില്ലാ എക്സിക്യൂട്ടീവ് ഓഫീസറാണ് ക്ഷേമനിധി അംഗങ്ങളുടെ കാര്യങ്ങൾ കൈകാര്യം ചെയ്യുന്നത്. ഈ ആനുകൂല്യത്തിന്റെ അപേക്ഷാ നടപടിക്രമം ഞങ്ങളുടെ ഉറവിടങ്ങളിൽ ലഭ്യമല്ല.",
                   source_id="SRC-PLNT-007", quote="Office of the District Executive Officer"))
    r = find(A, scheme_id=sid, type_id="spwwfb_office")
    if r:
        setv(A, r[0], "source_id", "SRC-PLNT-007")
        setv(A, r[0], "quote", "Office of the Chief Executive Officer, Azad Lane, Kottayam - 686 001")
for sid, did, src_ in [("PLNT-01", "spwwfb_form_14", "SRC-PLNT-008"), ("PLNT-04", "spwwfb_form_18", "SRC-PLNT-009")]:
    delete_where(SD, scheme_id=sid, doc_id=did)
    append(SD, dict(scheme_id=sid, doc_id=did, source_id=src_))
r = find(SC, scheme_id="PLNT-01")[0]
setv(SC, r, "notes", "Act 17 of 2008 s.3(5)(a). Form 14 (Scheme 2009, clause 38(1)) is the application; clause 38 itself (contribution period, amount) is not collected, so documents and procedure stay partial.")
r = find(SC, scheme_id="PLNT-04")[0]
setv(SC, r, "notes", "Act 17 of 2008 s.3(5)(d). Form 18 (Scheme 2009, clause 45(1)) is the application; clause 45 itself is not collected.")
delete_where(SD, scheme_id="PLNT-05", doc_id="aadhaar")
delete_where(SD, scheme_id="PLNT-05", doc_id="aadhaar_or_enrolment")
append(SD, dict(scheme_id="PLNT-05", doc_id="aadhaar_or_enrolment", source_id="SRC-PLNT-002",
                quote="Any individual eligible for receiving the benefits under the above Scheme shall hereby be required to furnish proof of possession of the Aadhaar number or undergo Aadhaar authentication"))
r = find(SC, scheme_id="PLNT-05")[0]
setv(SC, r, "notes", "G.O.(P) No.81/2024/LBR lists the Kerala Small Plantation Workers' Welfare Fund Board (item 10) among the 17 boards whose female members get the Maternity Benefit Scheme. Amount, number of deliveries and other conditions are in the 'extant Scheme guidelines', not collected.")
scheme("PLNT-06", name_en="Assistance to Dependants after a Small Plantation Worker Member's Death", name_ml="അംഗത്തിന്റെ മരണാനന്തരം ആശ്രിതർക്ക് ധനസഹായം",
       authority_en="Kerala Small Plantation Workers' Welfare Fund Board", authority_ml="കേരള ചെറുകിട തോട്ടം തൊഴിലാളി ക്ഷേമനിധി ബോർഡ്", category="plantation",
       summary_en="Financial assistance to dependants after a member's death, applied for with Form 20. The amount and eligibility are in Scheme clause 45, which we do not have yet.",
       summary_ml="അംഗം മരിച്ചാൽ ആശ്രിതർക്ക് ധനസഹായം; ഫാറം 20 വഴി അപേക്ഷിക്കണം. തുകയും അർഹതയും പദ്ധതി ഖണ്ഡിക 45-ൽ; അത് ഇതുവരെ ലഭ്യമല്ല.",
       eligibility_verification="informational", documents_verification="partial", apply_verification="partial",
       notes="Form 20 only. Needs Scheme 2009 clause 45.")
docs_for("PLNT-06", [("spwwfb_form_20", None)], source="SRC-PLNT-010")
apply_for("PLNT-06", [dict(scheme_id="PLNT-06", type_id="plantation_district_office", mode="in_person", note_en="Ask the District Executive Officer for your area.",
                           note_ml="നിങ്ങളുടെ പ്രദേശത്തെ ജില്ലാ എക്സിക്യൂട്ടീവ് ഓഫീസറോട് അന്വേഷിക്കുക.", source_id="SRC-PLNT-007", quote="Office of the District Executive Officer")])
scheme("PLNT-07", name_en="Family Pension (Small Plantation Workers' Welfare Fund)", name_ml="കുടുംബ പെൻഷൻ (ചെറുകിട തോട്ടം തൊഴിലാളി ക്ഷേമനിധി)",
       authority_en="Kerala Small Plantation Workers' Welfare Fund Board", authority_ml="കേരള ചെറുകിട തോട്ടം തൊഴിലാളി ക്ഷേമനിധി ബോർഡ്", category="plantation",
       summary_en="The Act allows the Fund to pay a family pension. Who qualifies and how much is set in the Scheme, which we do not have yet.",
       summary_ml="കുടുംബ പെൻഷൻ നൽകാൻ നിയമം അനുവദിക്കുന്നു. അർഹതയും തുകയും പദ്ധതിയിലാണ്; അത് ഇതുവരെ ലഭ്യമല്ല.",
       eligibility_verification="informational", documents_verification="partial", apply_verification="partial", notes="Act 17 of 2008 s.3(5)(b) only.")
apply_for("PLNT-07", [dict(scheme_id="PLNT-07", type_id="plantation_district_office", mode="in_person", note_en="Ask the District Executive Officer for your area.",
                           note_ml="നിങ്ങളുടെ പ്രദേശത്തെ ജില്ലാ എക്സിക്യൂട്ടീവ് ഓഫീസറോട് അന്വേഷിക്കുക.", source_id="SRC-PLNT-007", quote="Office of the District Executive Officer")])

# ---------------------------------------------------------------- Matsyafed offices
mf = [
    ("Thiruvananthapuram", "Muttathara, near Matsyafed Net Factory, Thiruvananthapuram", "9526041126"),
    ("Kollam", "Anila House, Sakthikulangara, Kollam", "0474-2772971, 9526041109"),
    ("Alappuzha", "Convent Square East Side, Alappuzha 688001", "0477-2241597, 9526041037"),
    ("Kottayam", "Vaikom, Kottayam 686141", "04829-216180, 9447232051"),
    ("Ernakulam", "Thoppumpady, Kochi, Ernakulam 682005", "0484-2222511, 9526041387"),
    ("Thrissur", "Mini Civil Station, Triprayar, P.O. Valapad, Thrissur", "0487-2396106, 9526041111"),
    ("Malappuram", "K.G. Padi, Tirur-1, Malappuram", "0494-2423503, 9526041060"),
    ("Kozhikode", "West Hill, Kozhikode", "0495-2380344, 9526041060"),
    ("Kannur", "Mopla Bay, Kannur", "0497-2731257, 9496243403"),
    ("Kasaragod", "Kasaba Beach, Kasaragod 671121", "04994-230176, 9526041127"),
]
ML = {'Thiruvananthapuram': 'തിരുവനന്തപുരം', 'Kollam': 'കൊല്ലം', 'Alappuzha': 'ആലപ്പുഴ', 'Kottayam': 'കോട്ടയം', 'Ernakulam': 'എറണാകുളം', 'Thrissur': 'തൃശ്ശൂർ',
      'Malappuram': 'മലപ്പുറം', 'Kozhikode': 'കോഴിക്കോട്', 'Kannur': 'കണ്ണൂർ', 'Kasaragod': 'കാസർകോട്'}
for i, (d, addr, ph) in enumerate(mf, 1):
    upsert(L, "location_id", dict(location_id=f"LOC-MFED-{i:03d}", type_id="matsyafed_office", name_en=f"Matsyafed District Office, {d}",
           name_ml=f"മത്സ്യഫെഡ് ജില്ലാ ഓഫീസ്, {ML[d]}", district=d, address=addr, phone=ph, source_id="SRC-MFED-001"))

# ---------------------------------------------------------------- terminology
for tid, en, ml, ctx, s in [
    ("matsyaboard", "Matsyaboard (Kerala Fishermen's Welfare Fund Board)", "മത്സ്യബോർഡ്", "Fishing welfare", K),
    ("fisheries_officer", "Fisheries Officer", "ഫിഷറീസ് ഓഫീസർ", "Where to apply", K),
    ("welfare_contribution", "Welfare Fund contribution", "ക്ഷേമനിധി വിഹിതം", "Fishing welfare", K),
    ("old_age_pension", "Old age pension", "വാർദ്ധക്യകാല പെൻഷൻ", "Fishing welfare", K),
    ("widow_pension", "Widow pension", "വിധവ പെൻഷൻ", "Fishing welfare", K),
    ("marriage_assistance", "Marriage assistance", "വിവാഹ ധനസഹായം", "Welfare", K),
    ("legal_heirship_certificate", "Legal heirship certificate", "ലീഗൽ ഹെയർഷിപ്പ് സർട്ടിഫിക്കറ്റ്", "Documents", K),
    ("small_plantation_worker", "Small plantation worker", "ചെറുകിട തോട്ടം തൊഴിലാളി", "Plantation welfare", "SRC-PLNT-001"),
]:
    upsert(T, "term_id", dict(term_id=tid, en=en, ml=ml, context=ctx, source_id=s))

# ---------------------------------------------------------------- profiles
def prof(pid, **kw):
    rows = find(PF, profile_id=pid)
    d = dict(profile_id=pid, **kw)
    if rows:
        for k, v in d.items():
            setv(PF, rows[0], k, v)
    else:
        append(PF, d)

for r in range(2, PF.max_row + 1):
    a = get(PF, r, "answers") or ""
    parts = [x.strip() for x in a.split(";") if x.strip() and not x.strip().startswith("fisher_beneficiary_status=")]
    setv(PF, r, "answers", "; ".join(parts))

exec(open(os.path.join(os.path.dirname(__file__), "family-event-2025.py"), encoding="utf-8").read())

exec(open(os.path.join(os.path.dirname(__file__), "profiles-2025.py"), encoding="utf-8").read())

wb.save(P)
print("saved")
