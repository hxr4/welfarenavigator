import openpyxl

P = "dataset/anavandi-dataset.xlsx"
wb = openpyxl.load_workbook(P)
F = wb["facts"]
h = [c.value for c in F[1]]
fid, q = h.index("fact_id") + 1, h.index("question_en") + 1

PLAIN = {
    "livelihood": "What work does your family do?",
    "registered_fisher": "Is the person applying a registered fisher?",
    "active_fisher": "Does the person applying still work as a fisher (an active fisher)?",
    "bpl": "Is your family in the BPL (below poverty line) category?",
    "bank_account": "Does the person applying have a bank account?",
    "traditional_fisher": "Is the person applying a registered traditional fisher?",
    "retired_from_fishing": "Has the person stopped fishing work (retired)?",
    "annual_family_income": "What is your family's total income in a year?",
    "kfwfb_membership_years": "How many years ago was the first Welfare Fund Board contribution paid?",
}
changed = []
for r in range(2, F.max_row + 1):
    k = F.cell(r, fid).value
    if k in PLAIN and F.cell(r, q).value != PLAIN[k]:
        changed.append((k, F.cell(r, q).value, PLAIN[k]))
        F.cell(r, q).value = PLAIN[k]
wb.save(P)
for c in changed:
    print(f"{c[0]}: {c[1]!r} -> {c[2]!r}")
